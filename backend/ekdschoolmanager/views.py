import csv
import io
import posixpath
import re
import zipfile
import xml.etree.ElementTree as ET
from datetime import date, datetime, time, timedelta
import base64
import binascii
import io
import re
import unicodedata
from urllib.parse import quote

from decimal import Decimal, InvalidOperation

from django.contrib.auth import authenticate
from django.db import transaction
from django.db.models import Count, F, Max, Prefetch, Q, Sum
from django.http import HttpResponse
from PIL import Image as PillowImage

from django.core.files.base import ContentFile
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.utils.formats import date_format
from django.utils.text import slugify
from rest_framework import serializers, status, viewsets
from rest_framework.authtoken.models import Token
from rest_framework.decorators import action
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.permissions import AllowAny, BasePermission
from rest_framework.response import Response
from rest_framework.views import APIView

from .models import CarriedDebt, CarriedDebtPayment, ReportCard, SessionClosure, ReportCardAppreciation, ReportCardSettings, SubjectCategoryOrder, AcademicSession, AcademicYear, Announcement, AnnouncementRead, AttendanceRecord, AttendanceSession, ClassFeeItem, Conversation, Message, MessageAttachment, ClassSubject, CustomUser, DisciplineRecord, ExpenseCategory, FeeInstallment, FeePayment, GradeEntry, GradeScheme, School, SchoolClass, SchoolExpense, SchoolLevel, SchoolMembership, StudentEnrollment, ClassGroupSession, ExcludedTimetableClass, Subject, SubjectCategory, SubjectPeriodRestriction, TeacherAssignmentSubject, TeacherClassAssignment, TeacherUnavailability, Timetable, TimetablePeriod, TuitionFeePlan
from .closures import ClosureError, close_session, closed_session_for_date
from .journey import search_students, student_journey
from .yearclosure import (
    YearClosureError, close_year, ensure_closable, final_decisions,
    outstanding_for, student_outcome,
)
from .assignment import (
    AVERAGE_BANDS, attach_fallback_averages, band_counts, band_limits,
    distribute_group, group_key, is_girl, natural_class_key, rules_for,
    stream_key, unassigned_students,
)
from .yearcopy import copy_year_settings
from .yearreopen import YearReopenError, reopen_year
from .timetable import DAY_LABELS, create_default_periods, describe_unplaced, irreducible_deficit, regenerate
from .timetable_pdf import class_timetables_pdf, teacher_timetables_pdf
from .reportcards import (
    SchemeCopyError, appreciations_for, copy_scheme, generate_class, generate_school,
    promotion_decision, role_holders, scheme_for, settings_for, signatories_for,
    signature_path, student_history, teacher_signatures_for, term_history,
)
from .reportcard_pdf import report_cards_pdf
from .serializers import AcademicSessionSerializer, AcademicYearSerializer, AnnouncementSerializer, AttendanceSessionSerializer, ConversationSerializer, MessageSerializer, CustomUserSerializer, DisciplineRecordSerializer, ExpenseCategorySerializer, FeePaymentSerializer, GradeSchemeSerializer, SchoolClassSerializer, SchoolExpenseSerializer, SchoolLevelSerializer, SchoolMembershipSerializer, SchoolSerializer, StudentEnrollmentSerializer, SubjectCategorySerializer, SubjectSerializer, TuitionFeePlanSerializer, normalize_togolese_phone


XLSX_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
OFFICE_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def read_xlsx_without_dependency(uploaded_file):
    """Lit la première feuille d'un XLSX quand openpyxl n'est pas installé."""
    with zipfile.ZipFile(uploaded_file) as archive:
        shared_strings = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            shared_strings = ["".join(node.text or "" for node in item.iter(f"{XLSX_NS}t")) for item in root]

        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        first_sheet = workbook.find(f".//{XLSX_NS}sheet")
        if first_sheet is None:
            return []
        relation_id = first_sheet.attrib.get(f"{OFFICE_REL_NS}id")
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        target = next(
            (relation.attrib["Target"] for relation in relationships.findall(f"{REL_NS}Relationship") if relation.attrib.get("Id") == relation_id),
            "worksheets/sheet1.xml",
        )
        sheet_path = posixpath.normpath(posixpath.join("xl", target))

        date_styles = set()
        if "xl/styles.xml" in archive.namelist():
            styles = ET.fromstring(archive.read("xl/styles.xml"))
            custom_formats = {
                int(item.attrib["numFmtId"]): item.attrib.get("formatCode", "").lower()
                for item in styles.findall(f".//{XLSX_NS}numFmt")
            }
            cell_formats = styles.find(f"{XLSX_NS}cellXfs")
            if cell_formats is not None:
                for index, style in enumerate(cell_formats):
                    format_id = int(style.attrib.get("numFmtId", 0))
                    format_code = custom_formats.get(format_id, "")
                    if format_id in range(14, 23) or all(marker in format_code for marker in ("y", "d")):
                        date_styles.add(index)

        sheet = ET.fromstring(archive.read(sheet_path))
        parsed_rows = []
        for row in sheet.findall(f".//{XLSX_NS}row"):
            values = {}
            for cell in row.findall(f"{XLSX_NS}c"):
                reference = cell.attrib.get("r", "A1")
                letters = re.match(r"[A-Z]+", reference)
                column = 0
                for letter in letters.group(0) if letters else "A":
                    column = column * 26 + ord(letter) - 64
                cell_type = cell.attrib.get("t")
                value_node = cell.find(f"{XLSX_NS}v")
                raw_value = value_node.text if value_node is not None else ""
                if cell_type == "inlineStr":
                    value = "".join(node.text or "" for node in cell.iter(f"{XLSX_NS}t"))
                elif cell_type == "s" and raw_value:
                    value = shared_strings[int(raw_value)]
                elif cell_type == "b":
                    value = raw_value == "1"
                elif raw_value and int(cell.attrib.get("s", 0)) in date_styles:
                    value = date(1899, 12, 30) + timedelta(days=float(raw_value))
                else:
                    value = raw_value
                values[column] = value
            parsed_rows.append([values.get(index, "") for index in range(1, max(values, default=0) + 1)])

    if not parsed_rows:
        return []
    headers = [str(value or "").strip() for value in parsed_rows[0]]
    return [dict(zip(headers, row)) for row in parsed_rows[1:] if any(value not in (None, "") for value in row)]


def accessible_schools(user):
    if user.is_superuser:
        return School.objects.filter(is_active=True)
    return School.objects.filter(
        Q(owner=user) | Q(memberships__user=user, memberships__is_active=True),
        is_active=True,
    ).distinct()


# ── Qui peut quoi dans un établissement ─────────────────────────────────────
#
# La direction dirige l'établissement : propriétaire, administrateur, censeur
# et proviseur y font les mêmes choses. Le censeur et le proviseur ne sont pas
# des adjoints à droits réduits — ils tiennent l'école quand le propriétaire
# n'est pas là, et n'ont donc rien de moins que lui.
DIRECTION_ROLES = [
    CustomUser.Role.OWNER,
    CustomUser.Role.ADMIN,
    CustomUser.Role.CENSEUR,
    CustomUser.Role.PROVISEUR,
]

# Le secrétariat s'y ajoute pour la vie scolaire — classes, inscriptions,
# années et sessions : c'est le travail quotidien du secrétaire et du
# surveillant. Le personnel, les matières, les finances et l'arrêt des comptes
# (clôture d'une session ou d'une année) restent à la direction.
SCHOOL_LIFE_ROLES = DIRECTION_ROLES + [
    CustomUser.Role.SECRETARY,
    CustomUser.Role.SURVEILLANT,
]


class LoginView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        username = str(request.data.get("username", "")).strip()
        user = authenticate(request, username=username, password=request.data.get("password", ""))
        if not user or user.is_archived:
            return Response({"message": "Nom d’utilisateur ou mot de passe incorrect."}, status=status.HTTP_400_BAD_REQUEST)

        token, _ = Token.objects.get_or_create(user=user)
        schools = accessible_schools(user)
        return Response({
            "token": token.key,
            "user": CustomUserSerializer(user, context={"request": request}).data,
            "schools": SchoolSerializer(
                schools,
                many=True,
                context={"request": request, "authenticated_user": user},
            ).data,
        })


class LogoutView(APIView):
    def post(self, request):
        Token.objects.filter(user=request.user).delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


class UsernameSuggestionView(APIView):
    def get(self, request):
        last_name = str(request.query_params.get("last_name", "")).strip()
        first_names = str(request.query_params.get("first_names", "")).strip()
        first_name = first_names.split()[0] if first_names else ""
        clean_last = slugify(last_name).replace("-", "")
        clean_first = slugify(first_name).replace("-", "")
        candidates = [
            clean_last,
            clean_first,
            f"{clean_last[:3]}{clean_first[:3]}",
            f"{clean_first[:3]}{clean_last[:3]}",
        ]
        candidates = list(dict.fromkeys(candidate for candidate in candidates if len(candidate) >= 3))
        if not candidates:
            return Response({"username": ""})
        username = next(
            (candidate for candidate in candidates if not CustomUser.objects.filter(username__iexact=candidate).exists()),
            None,
        )
        if username is None:
            base = candidates[2] if len(candidates) > 2 else candidates[0]
            suffix = 2
            username = f"{base}{suffix}"
            while CustomUser.objects.filter(username__iexact=username).exists():
                suffix += 1
                username = f"{base}{suffix}"
        return Response({"username": username})


class RoleChoicesView(APIView):
    def get(self, request):
        personnel_roles = {
            CustomUser.Role.ADMIN,
            CustomUser.Role.TEACHER,
            CustomUser.Role.ACCOUNTANT,
            CustomUser.Role.SECRETARY,
            CustomUser.Role.CENSEUR,
            CustomUser.Role.PROVISEUR,
            CustomUser.Role.SURVEILLANT,
            CustomUser.Role.STAFF,
        }
        roles = [
            {"value": value, "label": label}
            for value, label in CustomUser.Role.choices
            if value in personnel_roles
        ]
        return Response(roles)


class SchoolViewSet(viewsets.ModelViewSet):
    serializer_class = SchoolSerializer

    def get_queryset(self):
        return accessible_schools(self.request.user)

    def perform_create(self, serializer):
        if not (self.request.user.is_superuser or self.request.user.role == CustomUser.Role.OWNER):
            raise serializers.ValidationError({"permission": "Seuls un propriétaire ou un superutilisateur peuvent créer une école."})
        if self.request.user.is_superuser:
            owner = serializer.validated_data.get("owner")
            if not owner:
                raise serializers.ValidationError({"owner": "Sélectionnez le propriétaire de l’école."})
        else:
            owner = self.request.user
        school = serializer.save(owner=owner)
        if owner.role != CustomUser.Role.OWNER:
            owner.role = CustomUser.Role.OWNER
            owner.save(update_fields=["role"])
        SchoolMembership.objects.get_or_create(
            school=school,
            user=owner,
            defaults={"role": CustomUser.Role.OWNER},
        )

    def perform_update(self, serializer):
        if not self.request.user.is_superuser and serializer.instance.owner_id != self.request.user.id:
            raise serializers.ValidationError({"permission": "Seul le propriétaire peut modifier cette école."})
        serializer.save()

    def perform_destroy(self, instance):
        if not self.request.user.is_superuser and instance.owner_id != self.request.user.id:
            raise serializers.ValidationError({"permission": "Seul le propriétaire peut supprimer cette école."})
        instance.is_active = False
        instance.save(update_fields=["is_active"])

    @action(detail=True, methods=["get", "post"], url_path="members")
    def members(self, request, pk=None):
        school = self.get_object()
        if request.method == "GET":
            memberships = school.memberships.select_related("user")
            return Response(SchoolMembershipSerializer(memberships, many=True).data)
        if not request.user.is_superuser and school.owner_id != request.user.id:
            return Response({"detail": "Seul le propriétaire peut ajouter un membre."}, status=403)
        serializer = SchoolMembershipSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        serializer.save(school=school)
        return Response(serializer.data, status=201)


class SchoolScopedMixin:
    school = None

    def get_school(self):
        if self.school is not None:
            return self.school
        school_id = self.kwargs.get("school_pk") or self.request.headers.get("X-School-ID")
        if not school_id:
            raise serializers.ValidationError({"school": "Sélectionnez une école."})
        try:
            self.school = accessible_schools(self.request.user).get(pk=school_id)
        except (School.DoesNotExist, ValueError):
            raise serializers.ValidationError({"school": "Vous n’avez pas accès à cette école."})
        return self.school

    def has_role_for_school(self, school, roles):
        """L'utilisateur tient-il l'un de ces rôles dans cet établissement ?

        Le propriétaire de l'école passe toujours, qu'il ait ou non une
        adhésion : l'école est la sienne.
        """
        return (
            self.request.user.is_superuser
            or school.owner_id == self.request.user.id
            or school.memberships.filter(
                user=self.request.user, role__in=roles, is_active=True,
            ).exists()
        )

    def ensure_manager(self):
        school = self.get_school()
        self.ensure_manager_for_school(school)

    def ensure_manager_for_school(self, school):
        if not self.has_role_for_school(school, DIRECTION_ROLES):
            raise serializers.ValidationError({
                "permission": "Seule la direction de l’établissement peut faire cela.",
            })

    def ensure_school_life(self):
        """Classes, inscriptions, années et sessions : direction et secrétariat."""
        if not self.has_role_for_school(self.get_school(), SCHOOL_LIFE_ROLES):
            raise serializers.ValidationError({
                "permission": "Vous ne pouvez pas gérer la vie scolaire de cette école.",
            })


class FinanceMixin(SchoolScopedMixin):
    def ensure_fee_configurator(self):
        if not self.has_role_for_school(self.get_school(), DIRECTION_ROLES):
            raise serializers.ValidationError({"permission": "Seule la direction de l’établissement peut configurer l’écolage."})

    def ensure_finance_manager(self):
        allowed_roles = DIRECTION_ROLES + [CustomUser.Role.ACCOUNTANT]
        if not self.has_role_for_school(self.get_school(), allowed_roles):
            raise serializers.ValidationError({"permission": "Vous ne pouvez pas gérer les finances de cette école."})

    def get_academic_year(self):
        year_id = self.request.headers.get("X-Academic-Year-ID")
        if not year_id:
            raise serializers.ValidationError({"academic_year": "Sélectionnez une année académique."})
        try:
            return AcademicYear.objects.get(pk=year_id, school=self.get_school())
        except (AcademicYear.DoesNotExist, ValueError):
            raise serializers.ValidationError({"academic_year": "Année académique invalide."})


class TuitionFeePlanListView(FinanceMixin, APIView):
    def get(self, request, school_pk):
        year = self.get_academic_year()
        plans = TuitionFeePlan.objects.filter(school=self.get_school(), academic_year=year).select_related(
            "school_class", "school_class__level"
        ).prefetch_related("items__fee_module", "items__installments")
        return Response(TuitionFeePlanSerializer(plans, many=True).data)

    def post(self, request, school_pk):
        self.ensure_fee_configurator()
        year = self.get_academic_year()
        serializer = TuitionFeePlanSerializer(data=request.data, context={"school": self.get_school(), "academic_year": year})
        serializer.is_valid(raise_exception=True)
        serializer.save(school=self.get_school(), academic_year=year)
        return Response(serializer.data, status=201)


class TuitionFeePlanDetailView(FinanceMixin, APIView):
    def get_object(self, pk):
        try:
            return TuitionFeePlan.objects.prefetch_related("items__fee_module", "items__installments").get(
                pk=pk, school=self.get_school(), academic_year=self.get_academic_year()
            )
        except TuitionFeePlan.DoesNotExist:
            raise serializers.ValidationError({"plan": "Barème introuvable."})

    def patch(self, request, school_pk, pk):
        self.ensure_fee_configurator()
        plan = self.get_object(pk)
        serializer = TuitionFeePlanSerializer(plan, data=request.data, partial=True, context={"school": self.get_school(), "academic_year": self.get_academic_year()})
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(serializer.data)

    def delete(self, request, school_pk, pk):
        self.ensure_fee_configurator()
        plan = self.get_object(pk)
        if plan.items.filter(payments__isnull=False).exists():
            return Response({"detail": "Ce barème possède des paiements et ne peut pas être supprimé."}, status=400)
        plan.delete()
        return Response(status=204)

    def post(self, request, school_pk, pk):
        self.ensure_fee_configurator()
        source = self.get_object(pk)
        target_ids = request.data.get("target_classes", [])
        if not isinstance(target_ids, list) or not target_ids:
            raise serializers.ValidationError({"target_classes": "Sélectionnez au moins une classe cible."})
        try:
            normalized_ids = {int(value) for value in target_ids}
        except (TypeError, ValueError):
            raise serializers.ValidationError({"target_classes": "Une ou plusieurs classes cibles sont invalides."})
        if source.school_class_id in normalized_ids:
            raise serializers.ValidationError({"target_classes": "La classe source ne peut pas être une cible."})
        target_classes = list(SchoolClass.objects.filter(
            pk__in=normalized_ids, school=self.get_school(), academic_year=self.get_academic_year(),
        ))
        if len(target_classes) != len(normalized_ids):
            raise serializers.ValidationError({"target_classes": "Une ou plusieurs classes cibles sont invalides."})
        existing_plans = {
            plan.school_class_id: plan for plan in TuitionFeePlan.objects.filter(
                school_class_id__in=normalized_ids,
            ).prefetch_related("items__payments")
        }
        blocked = [school_class.name for school_class in target_classes if (
            existing_plans.get(school_class.id)
            and existing_plans[school_class.id].items.filter(payments__isnull=False).exists()
        )]
        if blocked:
            return Response({"detail": f"Copie impossible : des paiements existent déjà pour {', '.join(blocked)}."}, status=400)
        source_items = list(source.items.select_related("fee_module").prefetch_related("installments"))
        created_plans = []
        with transaction.atomic():
            for target_class in target_classes:
                existing = existing_plans.get(target_class.id)
                if existing:
                    existing.delete()
                target = TuitionFeePlan.objects.create(
                    school=self.get_school(), academic_year=self.get_academic_year(), school_class=target_class,
                )
                for source_item in source_items:
                    target_item = ClassFeeItem.objects.create(
                        plan=target, fee_module=source_item.fee_module,
                        male_amount=source_item.male_amount, female_amount=source_item.female_amount,
                        payable_in_installments=source_item.payable_in_installments,
                    )
                    FeeInstallment.objects.bulk_create([
                        FeeInstallment(
                            class_fee=target_item, name=row.name, percentage=row.percentage,
                            due_date=row.due_date, order=row.order,
                        ) for row in source_item.installments.all()
                    ])
                created_plans.append(target.id)
        targets = TuitionFeePlan.objects.filter(pk__in=created_plans).prefetch_related("items__fee_module", "items__installments")
        return Response(TuitionFeePlanSerializer(targets, many=True).data, status=201)


class CarriedDebtListView(FinanceMixin, APIView):
    """Écolages restés impayés à la clôture d'une année précédente.

    Ils ne dépendent plus d'une inscription : la dette suit l'élève et se
    règle pendant n'importe quelle année, d'où une liste à part de la collecte
    courante.
    """

    def serialize(self, debt):
        return {
            "id": debt.id,
            "student": debt.student.get_full_name(),
            "student_id": debt.student_id,
            "origin_year": debt.origin_year.name,
            "matricule": debt.origin_enrollment.enrollment_number,
            "amount": str(debt.amount),
            "settled_amount": str(debt.settled_amount),
            "outstanding": str(debt.outstanding),
            "payments": [
                {
                    "id": payment.id,
                    "amount": str(payment.amount),
                    "paid_on": payment.paid_on,
                    "method": payment.method,
                    "reference": payment.reference,
                    "received_by": payment.received_by.get_full_name(),
                }
                for payment in debt.payments.all()
            ],
        }

    def get_queryset(self):
        queryset = CarriedDebt.objects.filter(
            school=self.get_school(),
        ).select_related("student", "origin_year", "origin_enrollment").prefetch_related(
            "payments__received_by",
        )
        if self.request.query_params.get("student"):
            queryset = queryset.filter(student_id=self.request.query_params["student"])
        # Par défaut on ne montre que ce qui reste à encaisser : une dette
        # soldée n'a plus rien à faire dans une liste de recouvrement.
        if self.request.query_params.get("settled") != "1":
            queryset = queryset.filter(settled_amount__lt=F("amount"))
        return queryset

    def get(self, request, school_pk):
        debts = list(self.get_queryset())
        return Response({
            "debts": [self.serialize(debt) for debt in debts],
            "outstanding_total": str(sum((debt.outstanding for debt in debts), Decimal("0"))),
        })

    def post(self, request, school_pk):
        """Encaisse un versement sur une dette reportée."""
        self.ensure_finance_manager()
        try:
            debt = CarriedDebt.objects.get(
                pk=request.data.get("debt"), school=self.get_school(),
            )
        except (CarriedDebt.DoesNotExist, TypeError, ValueError):
            raise serializers.ValidationError({"debt": "Dette introuvable."})

        try:
            amount = Decimal(str(request.data.get("amount")))
        except (InvalidOperation, TypeError, ValueError):
            raise serializers.ValidationError({"amount": "Montant invalide."})
        if amount <= 0:
            raise serializers.ValidationError({"amount": "Le versement doit être positif."})
        if amount > debt.outstanding:
            raise serializers.ValidationError({"amount": (
                f"Il ne reste que {debt.outstanding} à régler sur cette dette."
            )})

        paid_on = parse_date(str(request.data.get("paid_on") or "")) or timezone.localdate()
        with transaction.atomic():
            CarriedDebtPayment.objects.create(
                debt=debt, amount=amount, paid_on=paid_on,
                method=str(request.data.get("method") or "espece"),
                reference=str(request.data.get("reference") or "")[:100],
                notes=str(request.data.get("notes") or "")[:2000],
                received_by=request.user,
            )
            debt.settled_amount += amount
            debt.save(update_fields=["settled_amount"])

        debt.refresh_from_db()
        return Response(self.serialize(debt))


class TuitionPaymentListView(FinanceMixin, APIView):
    def get(self, request, school_pk):
        year = self.get_academic_year()
        queryset = FeePayment.objects.filter(
            enrollment__school=self.get_school(), enrollment__academic_year=year,
        ).select_related("enrollment__student", "enrollment__school_class", "class_fee__fee_module", "installment", "received_by")
        if request.query_params.get("class_id"):
            queryset = queryset.filter(enrollment__school_class_id=request.query_params["class_id"])
        return Response(FeePaymentSerializer(queryset, many=True).data)

    def post(self, request, school_pk):
        self.ensure_finance_manager()
        serializer = FeePaymentSerializer(data=request.data, context={"school": self.get_school(), "academic_year": self.get_academic_year()})
        serializer.is_valid(raise_exception=True)
        serializer.save(received_by=request.user)
        return Response(serializer.data, status=201)


class TuitionComplianceView(FinanceMixin, APIView):
    def get(self, request, school_pk):
        year = self.get_academic_year()
        class_id = request.query_params.get("class_id")
        target = request.query_params.get("target")
        if not class_id:
            raise serializers.ValidationError({"class_id": "Sélectionnez une classe."})
        try:
            plan = TuitionFeePlan.objects.prefetch_related("items__fee_module", "items__installments").get(
                school=self.get_school(), academic_year=year, school_class_id=class_id,
            )
        except TuitionFeePlan.DoesNotExist:
            raise serializers.ValidationError({"plan": "Définissez d’abord le barème de cette classe."})
        if not target:
            raise serializers.ValidationError({"target": "Sélectionnez une rubrique de frais."})
        installment = None
        try:
            if target.startswith("installment:"):
                installment = FeeInstallment.objects.select_related("class_fee__fee_module").get(pk=target.split(":", 1)[1], class_fee__plan=plan)
                class_fee = installment.class_fee
            else:
                class_fee_id = target.split(":", 1)[1] if target.startswith("item:") else target
                class_fee = ClassFeeItem.objects.select_related("fee_module").get(pk=class_fee_id, plan=plan)
        except (ValueError, IndexError, FeeInstallment.DoesNotExist, ClassFeeItem.DoesNotExist):
            raise serializers.ValidationError({"target": "Rubrique ou tranche invalide."})
        enrollments = StudentEnrollment.objects.filter(
            school=self.get_school(), academic_year=year, school_class_id=class_id, status=StudentEnrollment.Status.ACTIVE,
        ).select_related("student", "school_class")
        rows = []
        for enrollment in enrollments:
            base_amount = class_fee.female_amount if enrollment.student.gender == CustomUser.Gender.FEMALE else class_fee.male_amount
            expected = base_amount if installment is None else (base_amount * installment.percentage / Decimal("100")).quantize(Decimal("0.01"))
            payments = enrollment.fee_payments.filter(class_fee=class_fee)
            if installment is not None:
                payments = payments.filter(installment=installment)
            paid = payments.aggregate(total=Sum("amount"))["total"] or Decimal("0")
            rows.append({
                "enrollment": enrollment.id, "enrollment_number": enrollment.enrollment_number,
                "student_name": enrollment.student.get_full_name(), "gender": enrollment.student.get_gender_display(),
                "student_status": enrollment.student.get_student_status_display(), "expected": expected,
                "paid": paid, "balance": max(expected - paid, Decimal("0")), "is_compliant": paid >= expected,
            })
        return Response({
            "plan": TuitionFeePlanSerializer(plan).data,
            "target": class_fee.fee_module.name if installment is None else f"{class_fee.fee_module.name} — {installment.name}",
            "compliant_count": sum(row["is_compliant"] for row in rows),
            "non_compliant_count": sum(not row["is_compliant"] for row in rows),
            "students": rows,
        })


class ExpenseCategoryListView(FinanceMixin, APIView):
    def get(self, request, school_pk):
        categories = ExpenseCategory.objects.filter(school=self.get_school(), is_active=True)
        return Response(ExpenseCategorySerializer(categories, many=True).data)

    def post(self, request, school_pk):
        self.ensure_finance_manager()
        serializer = ExpenseCategorySerializer(data=request.data, context={"school": self.get_school()})
        serializer.is_valid(raise_exception=True)
        serializer.save(school=self.get_school())
        return Response(serializer.data, status=201)


class SchoolExpenseListView(FinanceMixin, APIView):
    def get(self, request, school_pk):
        year = self.get_academic_year()
        queryset = SchoolExpense.objects.filter(
            school=self.get_school(), academic_year=year,
        ).select_related("category", "recorded_by")
        category_id = request.query_params.get("category")
        if category_id:
            queryset = queryset.filter(category_id=category_id)
        if request.query_params.get("date_from"):
            queryset = queryset.filter(expense_date__gte=request.query_params["date_from"])
        if request.query_params.get("date_to"):
            queryset = queryset.filter(expense_date__lte=request.query_params["date_to"])
        search = request.query_params.get("search", "").strip()
        if search:
            queryset = queryset.filter(Q(label__icontains=search) | Q(beneficiary__icontains=search) | Q(reference__icontains=search))
        return Response(SchoolExpenseSerializer(queryset, many=True).data)

    def post(self, request, school_pk):
        self.ensure_finance_manager()
        serializer = SchoolExpenseSerializer(data=request.data, context={"school": self.get_school()})
        serializer.is_valid(raise_exception=True)
        serializer.save(school=self.get_school(), academic_year=self.get_academic_year(), recorded_by=request.user)
        return Response(serializer.data, status=201)


class SchoolExpenseDetailView(FinanceMixin, APIView):
    def get_object(self, pk):
        try:
            return SchoolExpense.objects.select_related("category", "recorded_by").get(
                pk=pk, school=self.get_school(), academic_year=self.get_academic_year(),
            )
        except SchoolExpense.DoesNotExist:
            raise serializers.ValidationError({"expense": "Dépense introuvable."})

    def patch(self, request, school_pk, pk):
        self.ensure_finance_manager()
        serializer = SchoolExpenseSerializer(
            self.get_object(pk), data=request.data, partial=True, context={"school": self.get_school()},
        )
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(serializer.data)

    def delete(self, request, school_pk, pk):
        self.ensure_finance_manager()
        self.get_object(pk).delete()
        return Response(status=204)


class DisciplineMixin(SchoolScopedMixin):
    def get_academic_year(self):
        year_id = self.request.headers.get("X-Academic-Year-ID")
        if not year_id:
            raise serializers.ValidationError({"academic_year": "Sélectionnez une année académique."})
        try:
            return AcademicYear.objects.get(pk=year_id, school=self.get_school())
        except (AcademicYear.DoesNotExist, ValueError):
            raise serializers.ValidationError({"academic_year": "Année académique invalide."})

    def ensure_discipline_access(self):
        school = self.get_school()
        allowed_roles = [
            CustomUser.Role.ADMIN, CustomUser.Role.OWNER, CustomUser.Role.CENSEUR,
            CustomUser.Role.PROVISEUR, CustomUser.Role.SURVEILLANT,
            CustomUser.Role.TEACHER, CustomUser.Role.SECRETARY,
        ]
        if not (
            self.request.user.is_superuser
            or school.owner_id == self.request.user.id
            or school.memberships.filter(user=self.request.user, role__in=allowed_roles, is_active=True).exists()
        ):
            raise serializers.ValidationError({"permission": "Vous ne pouvez pas accéder à la discipline de cette école."})


class DisciplineRecordListView(DisciplineMixin, APIView):
    def get(self, request, school_pk):
        self.ensure_discipline_access()
        year = self.get_academic_year()
        enrollments = StudentEnrollment.objects.filter(
            school=self.get_school(), academic_year=year, status=StudentEnrollment.Status.ACTIVE,
        ).select_related("student", "school_class", "level").order_by("student__last_name", "student__first_name")
        records = DisciplineRecord.objects.filter(school=self.get_school(), academic_year=year).select_related(
            "enrollment__student", "enrollment__school_class", "recorded_by",
        )
        enrollment_id = request.query_params.get("enrollment")
        if enrollment_id:
            records = records.filter(enrollment_id=enrollment_id)
        entry_type = request.query_params.get("entry_type")
        if entry_type in {DisciplineRecord.EntryType.LATE, DisciplineRecord.EntryType.ABSENCE, DisciplineRecord.EntryType.INCIDENT}:
            records = records.filter(entry_type=entry_type)
        total_late_hours = records.filter(entry_type=DisciplineRecord.EntryType.LATE).aggregate(total=Sum("late_hours"))["total"] or Decimal("0")
        total_absence_hours = records.filter(entry_type=DisciplineRecord.EntryType.ABSENCE).aggregate(total=Sum("late_hours"))["total"] or Decimal("0")
        return Response({
            "students": [{
                "id": item.id,
                "enrollment_number": item.enrollment_number,
                "student_name": item.student.get_full_name(),
                "class_name": item.school_class.name if item.school_class else "Sans classe",
                "level_name": item.level.name if item.level else "",
            } for item in enrollments],
            "records": DisciplineRecordSerializer(records, many=True).data,
            "summary": {
                "total_late_hours": total_late_hours,
                "total_absence_hours": total_absence_hours,
                "absence_count": records.filter(entry_type=DisciplineRecord.EntryType.ABSENCE).count(),
                "incident_count": records.filter(entry_type=DisciplineRecord.EntryType.INCIDENT).count(),
                "record_count": records.count(),
            },
        })

    def post(self, request, school_pk):
        self.ensure_discipline_access()
        year = self.get_academic_year()
        serializer = DisciplineRecordSerializer(data=request.data, context={"school": self.get_school(), "academic_year": year})
        serializer.is_valid(raise_exception=True)
        # La discipline est datée, pas rattachée à une session : c'est la date
        # de l'incident qui dit si la période est déjà arrêtée.
        enrollment = serializer.validated_data.get("enrollment")
        closed = closed_session_for_date(
            year,
            enrollment.school_class if enrollment else None,
            serializer.validated_data.get("occurred_on"),
        )
        if closed is not None:
            raise serializers.ValidationError(
                {"occurred_on": f"La session « {closed.name} » est clôturée : cette date n’accepte plus de saisie."}
            )
        serializer.save(school=self.get_school(), academic_year=year, recorded_by=request.user)
        return Response(serializer.data, status=201)


class AttendanceMixin(SchoolScopedMixin):
    """Accès à l'appel : ceux qui font la classe et ceux qui la surveillent."""

    def get_academic_year(self):
        year_id = self.request.headers.get("X-Academic-Year-ID")
        if not year_id:
            raise serializers.ValidationError({"academic_year": "Sélectionnez une année académique."})
        try:
            return AcademicYear.objects.get(pk=year_id, school=self.get_school())
        except (AcademicYear.DoesNotExist, ValueError):
            raise serializers.ValidationError({"academic_year": "Année académique invalide."})

    def ensure_attendance_access(self):
        school = self.get_school()
        allowed_roles = [
            CustomUser.Role.ADMIN, CustomUser.Role.OWNER, CustomUser.Role.CENSEUR,
            CustomUser.Role.PROVISEUR, CustomUser.Role.SURVEILLANT,
            CustomUser.Role.TEACHER, CustomUser.Role.SECRETARY,
        ]
        if not (
            self.request.user.is_superuser
            or school.owner_id == self.request.user.id
            or school.memberships.filter(
                user=self.request.user, role__in=allowed_roles, is_active=True,
            ).exists()
        ):
            raise serializers.ValidationError(
                {"permission": "Vous ne pouvez pas accéder à l’appel de cette école."}
            )

    def get_class(self, year, class_id):
        try:
            return SchoolClass.objects.select_related("level").get(
                pk=class_id, school=self.get_school(), academic_year=year,
            )
        except (SchoolClass.DoesNotExist, ValueError, TypeError):
            raise serializers.ValidationError({"school_class": "Classe invalide."})


class AttendanceSessionListView(AttendanceMixin, APIView):
    """Appels d'une classe et relevé du jour."""

    def get(self, request, school_pk):
        self.ensure_attendance_access()
        year = self.get_academic_year()
        sessions = AttendanceSession.objects.filter(
            school=self.get_school(), academic_year=year,
        ).select_related("school_class", "class_subject__subject", "taken_by").prefetch_related(
            "records__enrollment__student",
        )

        class_id = request.query_params.get("school_class")
        if class_id:
            sessions = sessions.filter(school_class_id=class_id)
        taken_on = request.query_params.get("taken_on")
        if taken_on:
            sessions = sessions.filter(taken_on=taken_on)

        # Un enseignant ne voit que les classes qui lui sont confiées : l'appel
        # d'une classe voisine ne le regarde pas.
        if not self.can_supervise():
            sessions = sessions.filter(
                school_class__in=self.taught_class_ids(year),
            )

        return Response({
            "sessions": AttendanceSessionSerializer(sessions[:200], many=True).data,
            "can_supervise": self.can_supervise(),
        })

    def can_supervise(self):
        """Vue d'ensemble : direction et surveillance, pas les enseignants."""
        school = self.get_school()
        roles = [
            CustomUser.Role.ADMIN, CustomUser.Role.OWNER, CustomUser.Role.CENSEUR,
            CustomUser.Role.PROVISEUR, CustomUser.Role.SURVEILLANT, CustomUser.Role.SECRETARY,
        ]
        return (
            self.request.user.is_superuser
            or school.owner_id == self.request.user.id
            or school.memberships.filter(
                user=self.request.user, role__in=roles, is_active=True,
            ).exists()
        )

    def taught_class_ids(self, year):
        return TeacherClassAssignment.objects.filter(
            teacher=self.request.user, school=self.get_school(), academic_year=year,
        ).values_list("school_class_id", flat=True)

    @transaction.atomic
    def post(self, request, school_pk):
        """Enregistre un appel. Refaire l'appel du même créneau le corrige."""
        self.ensure_attendance_access()
        year = self.get_academic_year()
        school_class = self.get_class(year, request.data.get("school_class"))

        if not self.can_supervise() and school_class.id not in set(self.taught_class_ids(year)):
            raise serializers.ValidationError(
                {"school_class": "Cette classe ne vous est pas confiée."}
            )

        taken_on = request.data.get("taken_on") or timezone.localdate()
        closed = closed_session_for_date(year, school_class, taken_on)
        if closed is not None:
            raise serializers.ValidationError(
                {"taken_on": f"La session « {closed.name} » est clôturée : l’appel de cette date n’est plus modifiable."}
            )
        class_subject_id = request.data.get("class_subject") or None
        if class_subject_id:
            if not ClassSubject.objects.filter(
                pk=class_subject_id, school_class=school_class,
            ).exists():
                raise serializers.ValidationError(
                    {"class_subject": "Cette matière n’est pas enseignée dans cette classe."}
                )

        session, _ = AttendanceSession.objects.update_or_create(
            school_class=school_class, taken_on=taken_on,
            class_subject_id=class_subject_id,
            period=str(request.data.get("period", "")).strip()[:40],
            defaults={
                "school": self.get_school(),
                "academic_year": year,
                "taken_by": request.user,
                "note": str(request.data.get("note", "")).strip(),
            },
        )

        entries = request.data.get("records")
        if entries is not None:
            if not isinstance(entries, list):
                raise serializers.ValidationError(
                    {"records": "Les présences doivent être envoyées sous forme de liste."}
                )
            # Les inscriptions valides sont celles de la classe : une ligne
            # portant un élève d'ailleurs est refusée plutôt qu'ignorée.
            allowed = set(
                StudentEnrollment.objects.filter(
                    school_class=school_class, status=StudentEnrollment.Status.ACTIVE,
                ).values_list("id", flat=True)
            )
            statuses = dict(AttendanceRecord.Status.choices)
            rows = []
            for entry in entries:
                enrollment_id = entry.get("enrollment")
                if enrollment_id not in allowed:
                    raise serializers.ValidationError(
                        {"records": "Un élève de la liste n’appartient pas à cette classe."}
                    )
                status_value = str(entry.get("status", AttendanceRecord.Status.PRESENT))
                if status_value not in statuses:
                    raise serializers.ValidationError({"records": "Statut de présence inconnu."})
                rows.append(AttendanceRecord(
                    session=session, enrollment_id=enrollment_id, status=status_value,
                    minutes_late=max(0, int(entry.get("minutes_late") or 0)),
                    comment=str(entry.get("comment", "")).strip()[:200],
                ))
            # Réécriture complète : l'appel corrigé remplace le précédent, sans
            # laisser de ligne orpheline pour un élève retiré de la classe.
            session.records.all().delete()
            AttendanceRecord.objects.bulk_create(rows)

        session.refresh_from_db()
        return Response(AttendanceSessionSerializer(session).data, status=201)


class AttendanceSheetView(AttendanceMixin, APIView):
    """Feuille d'appel d'une classe pour une date : élèves et relevé existant."""

    def get(self, request, school_pk):
        self.ensure_attendance_access()
        year = self.get_academic_year()
        school_class = self.get_class(year, request.query_params.get("school_class"))
        taken_on = request.query_params.get("taken_on") or str(timezone.localdate())

        enrollments = StudentEnrollment.objects.filter(
            school_class=school_class, status=StudentEnrollment.Status.ACTIVE,
        ).select_related("student").order_by("student__last_name", "student__first_name")

        session = AttendanceSession.objects.filter(
            school_class=school_class, taken_on=taken_on,
            class_subject_id=request.query_params.get("class_subject") or None,
            period=request.query_params.get("period", ""),
        ).prefetch_related("records").first()
        existing = {record.enrollment_id: record for record in session.records.all()} if session else {}

        return Response({
            "school_class": {"id": school_class.id, "name": school_class.group},
            "taken_on": taken_on,
            # Vrai quand l'appel a déjà été fait : le mobile distingue alors
            # « aucun absent » de « appel non fait ».
            "is_taken": session is not None,
            "session_id": session.id if session else None,
            "students": [
                {
                    "enrollment": item.id,
                    "enrollment_number": item.enrollment_number,
                    "student_name": item.student.get_full_name(),
                    "status": existing[item.id].status if item.id in existing
                              else AttendanceRecord.Status.PRESENT,
                    "minutes_late": existing[item.id].minutes_late if item.id in existing else 0,
                    "comment": existing[item.id].comment if item.id in existing else "",
                }
                for item in enrollments
            ],
            "statuses": [
                {"value": value, "label": label}
                for value, label in AttendanceRecord.Status.choices
            ],
        })


class GradeMixin(SchoolScopedMixin):
    def get_academic_year(self):
        try:
            return AcademicYear.objects.get(pk=self.request.headers.get("X-Academic-Year-ID"), school=self.get_school())
        except (AcademicYear.DoesNotExist, ValueError, TypeError):
            raise serializers.ValidationError({"academic_year": "Sélectionnez une année académique valide."})

    def can_configure_grades(self):
        return self.has_role_for_school(self.get_school(), DIRECTION_ROLES)

    def ensure_grade_access(self):
        if self.can_configure_grades():
            return
        if self.request.user.role == CustomUser.Role.TEACHER or self.get_school().memberships.filter(user=self.request.user, role=CustomUser.Role.TEACHER, is_active=True).exists():
            return
        raise serializers.ValidationError({"permission": "Vous n’avez pas accès aux notes."})

    def ensure_grade_configurator(self):
        if not self.can_configure_grades():
            raise serializers.ValidationError({"permission": "Seule la direction de l’établissement peut configurer les notes."})

    def can_access_class_subject(self, class_subject):
        if self.can_configure_grades():
            return True
        return TeacherAssignmentSubject.objects.filter(
            assignment__teacher=self.request.user,
            assignment__school=self.get_school(),
            assignment__academic_year=class_subject.school_class.academic_year,
            assignment__school_class=class_subject.school_class,
            class_subject=class_subject,
        ).exists()


class GradeContextView(GradeMixin, APIView):
    def get(self, request, school_pk):
        self.ensure_grade_access()
        year = self.get_academic_year()
        today = timezone.localdate()
        # Une session close reste listée : elle est en lecture seule, pas
        # masquée. Le drapeau `is_closed` dit au client de désactiver la saisie.
        configured = set(
            GradeScheme.objects.filter(session__academic_year=year)
            .values_list("session_id", flat=True)
        )
        sessions = AcademicSession.objects.filter(academic_year=year).exclude(is_active=False, is_closed=False).prefetch_related("classes__level", "classes__subject_configurations__subject")
        result = []
        for session in sessions:
            classes = []
            for school_class in session.classes.all():
                subjects = [config for config in school_class.subject_configurations.all() if self.can_access_class_subject(config)]
                if subjects or self.can_configure_grades():
                    classes.append({
                        "id": school_class.id,
                        "name": " ".join(part for part in (school_class.level.name, school_class.series, school_class.group) if part),
                        "subjects": [{"id": config.id, "name": config.subject.name} for config in subjects],
                    })
            if classes:
                result.append({
                    "id": session.id, "name": session.name, "label": session.label,
                    # Les dates permettent de désigner la session en cours ;
                    # sans elles, un client ne peut que deviner laquelle
                    # afficher quand plusieurs sessions sont actives.
                    "start_date": session.start_date,
                    "end_date": session.end_date,
                    "is_current": session.start_date <= today <= session.end_date,
                    "is_closed": session.is_closed,
                    # Dernière session : c'est elle qui porte la moyenne
                    # annuelle et les décisions de passage.
                    "is_final": session.is_final,
                    "has_scheme": session.id in configured,
                    "classes": classes,
                })
        return Response({"can_configure": self.can_configure_grades(), "sessions": result})


class GradeSchemeView(GradeMixin, APIView):
    def get_session(self, session_pk):
        try:
            return AcademicSession.objects.get(pk=session_pk, academic_year=self.get_academic_year())
        except (AcademicSession.DoesNotExist, TypeError, ValueError):
            raise serializers.ValidationError({"session": "Session académique invalide."})

    def get(self, request, school_pk, session_pk):
        self.ensure_grade_access()
        session = self.get_session(session_pk)
        try:
            scheme = GradeScheme.objects.prefetch_related("lines", "groups__lines").get(session=session)
        except GradeScheme.DoesNotExist:
            return Response(None)
        return Response(GradeSchemeSerializer(scheme).data)

    def put(self, request, school_pk, session_pk):
        self.ensure_grade_configurator()
        session = self.get_session(session_pk)
        # Modifier un coefficient après clôture recalculerait des moyennes
        # déjà remises aux familles.
        if session.is_closed:
            raise serializers.ValidationError(
                {"session": "Cette session est clôturée : le barème n’est plus modifiable."}
            )
        scheme = GradeScheme.objects.filter(session=session).first()
        serializer = GradeSchemeSerializer(scheme, data=request.data) if scheme else GradeSchemeSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        serializer.save(session=session, created_by=request.user) if scheme is None else serializer.save()
        return Response(serializer.data, status=201 if scheme is None else 200)

    def get_source_session(self, session_pk):
        """Session dont on reprend la configuration, quelle que soit son année.

        Un établissement garde ses habitudes d'une année sur l'autre : le
        barème de l'an dernier est souvent le meilleur point de départ. La
        source se cherche donc dans toute l'école, pas seulement dans l'année
        active — seule la destination reste l'année en cours.
        """
        try:
            return AcademicSession.objects.select_related("academic_year").get(
                pk=session_pk, academic_year__school=self.get_school(),
            )
        except (AcademicSession.DoesNotExist, TypeError, ValueError):
            raise serializers.ValidationError({"source": "Session académique invalide."})

    def post(self, request, school_pk, session_pk):
        """Recopie ici la configuration d'une autre session, même d'une autre année."""
        self.ensure_grade_configurator()
        target = self.get_session(session_pk)
        source = self.get_source_session(request.data.get("source"))
        try:
            scheme = copy_scheme(source, target, user=request.user)
        except SchemeCopyError as issue:
            raise serializers.ValidationError({"source": str(issue)})
        return Response(GradeSchemeSerializer(scheme).data, status=201)


class GradeSchemeSourceView(GradeMixin, APIView):
    """Configurations de notes déjà en place, année par année.

    Alimente le choix « reprendre une configuration existante ». Les années
    sont rendues de la plus récente à la plus ancienne : c'est presque toujours
    la dernière que l'on veut reprendre.
    """

    def get(self, request, school_pk):
        self.ensure_grade_access()
        configured = (
            AcademicSession.objects
            .filter(
                academic_year__school=self.get_school(),
                grade_scheme__isnull=False,
            )
            .select_related("academic_year")
            .order_by("-academic_year__start_date", "start_date")
        )
        years = []
        for session in configured:
            year = session.academic_year
            if not years or years[-1]["id"] != year.id:
                years.append({
                    "id": year.id, "name": year.name,
                    "is_active": year.is_active, "sessions": [],
                })
            years[-1]["sessions"].append({
                "id": session.id, "name": session.name, "label": session.label,
                "is_closed": session.is_closed, "is_final": session.is_final,
            })
        return Response({"years": years})


class GradeSheetView(GradeMixin, APIView):
    def get_context(self, session_pk, class_subject_pk):
        year = self.get_academic_year()
        try:
            session = AcademicSession.objects.get(pk=session_pk, academic_year=year)
            class_subject = ClassSubject.objects.select_related("school_class", "subject").get(pk=class_subject_pk, school_class__academic_year=year, school_class__school=self.get_school())
        except (AcademicSession.DoesNotExist, ClassSubject.DoesNotExist):
            raise serializers.ValidationError({"context": "Session, classe ou matière invalide."})
        if not session.classes.filter(pk=class_subject.school_class_id).exists():
            raise serializers.ValidationError({"session": "Cette classe n’appartient pas à la session."})
        if not self.can_access_class_subject(class_subject):
            raise serializers.ValidationError({"permission": "Cette matière ne vous est pas affectée dans cette classe."})
        try:
            scheme = GradeScheme.objects.prefetch_related("lines", "groups__lines").get(session=session)
        except GradeScheme.DoesNotExist:
            raise serializers.ValidationError({"configuration": "Les lignes de notes ne sont pas encore configurées pour cette session."})
        return session, class_subject, scheme

    @staticmethod
    def calculate_average(scheme, scores):
        lines = list(scheme.lines.all())
        if not lines or any(line.id not in scores for line in lines):
            return None
        normalized = {line.id: scores[line.id] * Decimal("20") / line.max_score for line in lines}
        if scheme.calculation_method == GradeScheme.CalculationMethod.WEIGHTED:
            return sum(normalized[line.id] * line.weight / Decimal("100") for line in lines)
        if scheme.calculation_method == GradeScheme.CalculationMethod.GROUPS:
            group_averages = []
            groups = list(scheme.groups.all())
            legacy_group_weight = Decimal("100") / len(groups)
            for group in groups:
                group_lines = list(group.lines.all())
                legacy_line_weight = Decimal("100") / len(group_lines)
                group_average = sum(
                    normalized[line.id] * (line.weight or legacy_line_weight) / Decimal("100")
                    for line in group_lines
                )
                group_averages.append(group_average * (group.weight or legacy_group_weight) / Decimal("100"))
            return sum(group_averages)
        return sum(normalized.values()) / len(normalized)

    def get(self, request, school_pk, session_pk, class_subject_pk):
        self.ensure_grade_access()
        session, class_subject, scheme = self.get_context(session_pk, class_subject_pk)
        enrollments = StudentEnrollment.objects.filter(school_class=class_subject.school_class, status=StudentEnrollment.Status.ACTIVE).select_related("student").order_by("student__last_name", "student__first_name")
        entries = GradeEntry.objects.filter(class_subject=class_subject, line__scheme=scheme)
        entry_map = {(entry.enrollment_id, entry.line_id): entry.score for entry in entries}
        rows = []
        for enrollment in enrollments:
            scores = {line.id: entry_map[(enrollment.id, line.id)] for line in scheme.lines.all() if (enrollment.id, line.id) in entry_map}
            average = self.calculate_average(scheme, scores)
            rows.append({"enrollment": enrollment.id, "matricule": enrollment.enrollment_number, "student_name": enrollment.student.get_full_name(), "scores": {str(key): value for key, value in scores.items()}, "average": average.quantize(Decimal("0.01")) if average is not None else None})
        return Response({"scheme": GradeSchemeSerializer(scheme).data, "class_subject": {"id": class_subject.id, "subject": class_subject.subject.name}, "students": rows})

    def post(self, request, school_pk, session_pk, class_subject_pk):
        self.ensure_grade_access()
        session, class_subject, scheme = self.get_context(session_pk, class_subject_pk)
        # La lecture reste ouverte sur une session close, la saisie non : les
        # notes archivées à la clôture ne doivent plus bouger.
        if session.is_closed:
            raise serializers.ValidationError(
                {"session": "Cette session est clôturée : les notes ne sont plus modifiables."}
            )
        if not session.is_active:
            raise serializers.ValidationError(
                {"session": "Cette session n’est pas active : la saisie est fermée."}
            )
        lines = {line.id: line for line in scheme.lines.all()}
        enrollment_ids = set(StudentEnrollment.objects.filter(school_class=class_subject.school_class, status=StudentEnrollment.Status.ACTIVE).values_list("id", flat=True))
        grades = request.data.get("grades", [])
        errors = []
        prepared = []
        for index, item in enumerate(grades, start=1):
            try:
                enrollment_id, line_id = int(item["enrollment"]), int(item["line"])
                score = Decimal(str(item["score"]))
                if enrollment_id not in enrollment_ids or line_id not in lines or score < 0 or score > lines[line_id].max_score:
                    raise ValueError
                prepared.append((enrollment_id, line_id, score))
            except (KeyError, TypeError, ValueError, ArithmeticError):
                errors.append(f"Note invalide à la ligne {index}.")
        if errors:
            return Response({"grades": errors}, status=400)
        with transaction.atomic():
            for enrollment_id, line_id, score in prepared:
                GradeEntry.objects.update_or_create(
                    enrollment_id=enrollment_id, line_id=line_id, class_subject=class_subject,
                    defaults={"score": score, "entered_by": request.user},
                )
        return self.get(request, school_pk, session_pk, class_subject_pk)


class CustomUserViewSet(SchoolScopedMixin, viewsets.ModelViewSet):
    serializer_class = CustomUserSerializer

    def get_queryset(self):
        school = self.get_school()
        return CustomUser.objects.filter(
            school_memberships__school=school,
            school_memberships__is_active=True,
            is_archived=False,
        ).distinct()

    def perform_create(self, serializer):
        self.ensure_manager()
        with transaction.atomic():
            user = serializer.save()
            SchoolMembership.objects.create(school=self.get_school(), user=user, role=user.role)

    def perform_update(self, serializer):
        self.ensure_manager()
        serializer.save()

    def perform_destroy(self, instance):
        self.ensure_manager()
        SchoolMembership.objects.filter(school=self.get_school(), user=instance).update(is_active=False)


class TeacherViewSet(CustomUserViewSet):
    def get_selected_academic_year(self):
        try:
            return AcademicYear.objects.get(pk=self.request.headers.get("X-Academic-Year-ID"), school=self.get_school())
        except (AcademicYear.DoesNotExist, ValueError, TypeError):
            raise serializers.ValidationError({"academic_year": "Sélectionnez une année académique valide."})

    def get_queryset(self):
        school = self.get_school()
        return CustomUser.objects.filter(
            school_memberships__school=school,
            school_memberships__role__in=[
                CustomUser.Role.ADMIN,
                CustomUser.Role.TEACHER,
                CustomUser.Role.ACCOUNTANT,
                CustomUser.Role.SECRETARY,
                CustomUser.Role.CENSEUR,
                CustomUser.Role.PROVISEUR,
                CustomUser.Role.SURVEILLANT,
                CustomUser.Role.STAFF,
            ],
            school_memberships__is_active=True,
            is_archived=False,
        ).select_related("primary_subject").prefetch_related("subjects").distinct()

    def perform_create(self, serializer):
        role = serializer.validated_data.get("role", CustomUser.Role.TEACHER)
        if role == CustomUser.Role.OWNER and not self.request.user.is_superuser:
            raise serializers.ValidationError({"role": "Seul un superutilisateur peut créer un propriétaire."})
        selected_ids = serializer.validated_data.pop("school_ids", [])
        if not selected_ids:
            selected_ids = [self.get_school().id]
        schools = list(accessible_schools(self.request.user).filter(id__in=set(selected_ids)))
        if len(schools) != len(set(selected_ids)):
            raise serializers.ValidationError({"school_ids": "Une ou plusieurs écoles sont inaccessibles."})
        for school in schools:
            self.ensure_manager_for_school(school)
        with transaction.atomic():
            user = serializer.save(role=role)
            SchoolMembership.objects.bulk_create([
                SchoolMembership(school=school, user=user, role=role)
                for school in schools
            ])

    def perform_update(self, serializer):
        self.ensure_manager()
        selected_ids = serializer.validated_data.pop("school_ids", None)
        role = serializer.validated_data.get("role", serializer.instance.role)
        if role == CustomUser.Role.OWNER and not self.request.user.is_superuser:
            raise serializers.ValidationError({"role": "Seul un superutilisateur peut attribuer le rôle propriétaire."})
        schools = None
        if selected_ids is not None:
            if not selected_ids:
                raise serializers.ValidationError({"school_ids": "Sélectionnez au moins une école."})
            schools = list(accessible_schools(self.request.user).filter(id__in=set(selected_ids)))
            if len(schools) != len(set(selected_ids)):
                raise serializers.ValidationError({"school_ids": "Une ou plusieurs écoles sont inaccessibles."})
            for school in schools:
                self.ensure_manager_for_school(school)
        with transaction.atomic():
            user = serializer.save(role=role)
            if schools is not None:
                SchoolMembership.objects.filter(user=user).exclude(school__in=schools).update(is_active=False)
                for school in schools:
                    SchoolMembership.objects.update_or_create(
                        school=school,
                        user=user,
                        defaults={"role": role, "is_active": True},
                    )

    @action(detail=True, methods=["get", "post"], url_path="classes")
    def class_assignments(self, request, pk=None, school_pk=None):
        teacher = self.get_object()
        year = self.get_selected_academic_year()
        if teacher.role != CustomUser.Role.TEACHER:
            return Response({"teacher": "Les classes peuvent être assignées uniquement à un enseignant."}, status=400)
        if request.method == "GET":
            assignments = TeacherClassAssignment.objects.filter(
                school=self.get_school(), academic_year=year, teacher=teacher,
            ).prefetch_related("subject_links__class_subject")
            conflicts = TeacherAssignmentSubject.objects.filter(
                assignment__school=self.get_school(), assignment__academic_year=year,
            ).exclude(assignment__teacher=teacher).select_related("assignment__teacher", "class_subject")
            return Response({
                "assignments": [{
                    "class_id": assignment.school_class_id,
                    "subject_ids": [link.class_subject.subject_id for link in assignment.subject_links.all()],
                } for assignment in assignments],
                "unavailable_subjects": [{
                    "class_id": link.assignment.school_class_id,
                    "subject_id": link.class_subject.subject_id,
                    "teacher_name": link.assignment.teacher.get_full_name(),
                } for link in conflicts],
            })
        self.ensure_manager()
        requested = request.data.get("assignments", [])
        class_ids = [item.get("class_id") for item in requested]
        if len(class_ids) != len(set(class_ids)):
            return Response({"assignments": "Une classe ne peut apparaître qu’une seule fois."}, status=400)
        classes = list(SchoolClass.objects.filter(
            id__in=class_ids, school=self.get_school(), academic_year=year, is_active=True,
        ))
        if len(classes) != len(class_ids):
            return Response({"assignments": "Une ou plusieurs classes sont invalides."}, status=400)
        class_map = {school_class.id: school_class for school_class in classes}
        prepared = []
        selected_config_ids = []
        teacher_subject_ids = set(teacher.subjects.values_list("id", flat=True))
        for item in requested:
            subject_ids = list(dict.fromkeys(item.get("subject_ids", [])))
            if not subject_ids:
                return Response({"assignments": f"Sélectionnez au moins une matière pour {class_map[item['class_id']].name}."}, status=400)
            configs = list(ClassSubject.objects.filter(school_class_id=item["class_id"], subject_id__in=subject_ids))
            if len(configs) != len(subject_ids):
                return Response({"assignments": "Une matière sélectionnée n’est pas configurée dans cette classe."}, status=400)
            invalid_subjects = [config.subject.name for config in configs if config.subject_id not in teacher_subject_ids]
            if invalid_subjects:
                return Response({"assignments": f"Cet enseignant n’enseigne pas : {', '.join(invalid_subjects)}."}, status=400)
            prepared.append((class_map[item["class_id"]], configs))
            selected_config_ids.extend(config.id for config in configs)
        conflicts = TeacherAssignmentSubject.objects.filter(
            class_subject_id__in=selected_config_ids,
        ).exclude(assignment__teacher=teacher).select_related("assignment__teacher", "class_subject__subject").first()
        if conflicts:
            return Response({"assignments": f"{conflicts.class_subject.subject.name} est déjà enseignée par {conflicts.assignment.teacher.get_full_name()} dans cette classe."}, status=400)
        with transaction.atomic():
            TeacherClassAssignment.objects.filter(school=self.get_school(), academic_year=year, teacher=teacher).delete()
            links = []
            for school_class, configs in prepared:
                assignment = TeacherClassAssignment.objects.create(school=self.get_school(), academic_year=year, teacher=teacher, school_class=school_class)
                links.extend(TeacherAssignmentSubject(assignment=assignment, class_subject=config) for config in configs)
            TeacherAssignmentSubject.objects.bulk_create(links)
        return Response({"assignments": len(prepared)})

    @action(detail=True, methods=["get", "post"], url_path="unavailability")
    def unavailability(self, request, pk=None, school_pk=None):
        teacher = self.get_object()
        year = self.get_selected_academic_year()
        if teacher.role != CustomUser.Role.TEACHER:
            return Response({"teacher": "Les indisponibilités concernent uniquement les enseignants."}, status=400)
        queryset = TeacherUnavailability.objects.filter(school=self.get_school(), academic_year=year, teacher=teacher)
        if request.method == "GET":
            return Response([{
                "id": item.id, "day": item.day, "day_label": item.get_day_display(), "all_day": item.all_day,
                "start_time": item.start_time.strftime("%H:%M") if item.start_time else None,
                "end_time": item.end_time.strftime("%H:%M") if item.end_time else None,
            } for item in queryset])
        self.ensure_manager()
        slots = request.data.get("slots", [])
        normalized = []
        by_day = {}
        for index, slot in enumerate(slots, start=1):
            try:
                day = int(slot.get("day"))
                if day not in range(7): raise ValueError
                all_day = bool(slot.get("all_day"))
                start = end = None
                if not all_day:
                    start = datetime.strptime(str(slot.get("start_time")), "%H:%M").time()
                    end = datetime.strptime(str(slot.get("end_time")), "%H:%M").time()
                    if start >= end: raise ValueError
            except (TypeError, ValueError):
                return Response({"slots": f"Plage invalide à la ligne {index}."}, status=400)
            existing = by_day.setdefault(day, [])
            if existing and (all_day or any(item[0] for item in existing)):
                return Response({"slots": f"Les indisponibilités du jour {day + 1} se chevauchent."}, status=400)
            if not all_day and any(not (end <= item[1] or start >= item[2]) for item in existing):
                return Response({"slots": f"Les indisponibilités du jour {day + 1} se chevauchent."}, status=400)
            existing.append((all_day, start, end))
            normalized.append((day, all_day, start, end))
        with transaction.atomic():
            queryset.delete()
            TeacherUnavailability.objects.bulk_create([
                TeacherUnavailability(school=self.get_school(), academic_year=year, teacher=teacher, day=day, all_day=all_day, start_time=start, end_time=end)
                for day, all_day, start, end in normalized
            ])
        return Response({"slots": len(normalized)})


class SubjectCategoryViewSet(SchoolScopedMixin, viewsets.ModelViewSet):
    """Types de matières définis librement par l'établissement."""

    serializer_class = SubjectCategorySerializer

    def get_queryset(self):
        return SubjectCategory.objects.filter(school=self.get_school())

    def perform_create(self, serializer):
        self.ensure_manager()
        serializer.save(school=self.get_school())

    def perform_update(self, serializer):
        self.ensure_manager()
        serializer.save()

    def perform_destroy(self, instance):
        self.ensure_manager()
        # Les matières classées ici retrouvent simplement un type vide.
        instance.delete()


class SubjectViewSet(SchoolScopedMixin, viewsets.ModelViewSet):
    serializer_class = SubjectSerializer

    def get_queryset(self):
        return Subject.objects.filter(
            school=self.get_school(), is_active=True,
        ).select_related("category")

    def perform_create(self, serializer):
        self.ensure_manager()
        serializer.save(school=self.get_school())

    def perform_update(self, serializer):
        self.ensure_manager()
        serializer.save()

    def perform_destroy(self, instance):
        self.ensure_manager()
        instance.is_active = False
        instance.save(update_fields=["is_active"])


class AcademicYearViewSet(SchoolScopedMixin, viewsets.ModelViewSet):
    serializer_class = AcademicYearSerializer

    def get_queryset(self):
        return AcademicYear.objects.filter(school=self.get_school()).prefetch_related(
            "sessions__classes",
            Prefetch(
                "sessions__closure",
                queryset=SessionClosure.objects.select_related("closed_by"),
            ),
        )

    def perform_create(self, serializer):
        self.ensure_school_life()
        with transaction.atomic():
            if serializer.validated_data.get("is_active"):
                current = AcademicYear.objects.filter(
                    school=self.get_school(), is_active=True
                ).first()
                if current:
                    raise serializers.ValidationError({
                        "is_active": f"Clôturez d’abord l’année académique active « {current.name} ».",
                    })
            serializer.save(school=self.get_school())

    def perform_update(self, serializer):
        self.ensure_school_life()
        with transaction.atomic():
            if (
                serializer.instance.is_active
                and "is_active" in serializer.validated_data
                and not serializer.validated_data["is_active"]
            ):
                raise serializers.ValidationError({
                    "is_active": "Une année active doit être clôturée pour devenir inactive.",
                })
            if serializer.validated_data.get("is_active") and not serializer.instance.is_active:
                current = AcademicYear.objects.filter(
                    school=self.get_school(), is_active=True
                ).exclude(pk=serializer.instance.pk).first()
                if current:
                    raise serializers.ValidationError({
                        "is_active": f"Clôturez d’abord l’année académique active « {current.name} ».",
                    })
            serializer.save()

    def perform_destroy(self, instance):
        self.ensure_school_life()
        if instance.is_active:
            raise serializers.ValidationError({"year": "Une année active ne peut pas être supprimée."})
        instance.delete()

    @action(detail=True, methods=["post"])
    def close(self, request, pk=None, school_pk=None):
        """Clôture l'année : archive, reporte les impayés, fait passer les élèves.

        Deux refus possibles, chacun avec ce qu'il reste à faire : une session
        encore ouverte, ou l'année suivante pas encore créée.
        """
        self.ensure_manager()
        academic_year = self.get_object()
        try:
            closure = close_year(academic_year, user=request.user)
        except YearClosureError as issue:
            return Response({"detail": str(issue)}, status=400)

        academic_year.refresh_from_db()
        return Response({
            **self.get_serializer(academic_year).data,
            "closure": {
                "next_year": closure.next_year.name if closure.next_year else "",
                "enrollments": closure.enrollment_count,
                "promoted": closure.promoted_count,
                "repeated": closure.repeated_count,
                "graduated": closure.graduated_count,
                "unassigned": closure.unassigned_count,
                "undecided": closure.undecided_count,
                "carried_debt": str(closure.carried_debt_total),
                # Ce qui a été reconduit sur l'année neuve, pour que
                # l'utilisateur sache ce qu'il retrouvera en y basculant.
                "copied": closure.payload.get("copied", {}),
            },
        })

    @action(detail=True, methods=["post"])
    def reopen(self, request, pk=None, school_pk=None):
        """Annule la clôture : défait ce qu'elle avait fait, et rouvre l'année.

        Refusée si ce n'est pas la dernière clôture de l'établissement, ou si
        du travail a été fait depuis dans l'année suivante — le message dit
        alors quoi défaire d'abord.
        """
        self.ensure_manager()
        academic_year = self.get_object()
        try:
            report = reopen_year(academic_year)
        except YearReopenError as issue:
            return Response({"detail": str(issue)}, status=400)

        academic_year.refresh_from_db()
        return Response({
            **self.get_serializer(academic_year).data,
            "reopened": report,
        })

    @action(detail=True, methods=["get"], url_path="closure-preview")
    def closure_preview(self, request, pk=None, school_pk=None):
        """Ce que la clôture ferait, sans rien écrire.

        Permet de voir d'avance combien d'élèves passent, redoublent ou
        sortent — et surtout combien resteraient sans classe faute d'une
        classe correspondante dans l'année suivante.

        La reconduction des classes est jouée pour de vrai, dans une
        transaction annulée avant de répondre : sans elle, l'aperçu
        annoncerait tout l'établissement sans classe alors que la clôture
        affectera chacun.
        """
        self.ensure_manager()
        academic_year = self.get_object()
        try:
            target_year = ensure_closable(academic_year)
        except YearClosureError as issue:
            return Response({"ready": False, "detail": str(issue)})

        rows = []
        with transaction.atomic():
            copied = copy_year_settings(academic_year, target_year)
            decisions = final_decisions(academic_year)
            for enrollment in StudentEnrollment.objects.filter(
                academic_year=academic_year, status=StudentEnrollment.Status.ACTIVE,
            ).select_related("student", "school_class__level", "level").prefetch_related("fee_payments"):
                outcome = student_outcome(enrollment, decisions.get(enrollment.id), target_year)
                rows.append({
                    "student": enrollment.student.get_full_name(),
                    "class": enrollment.school_class.group if enrollment.school_class else "",
                    "outcome": outcome["kind"],
                    "detail": outcome["detail"],
                    "next_class": outcome["school_class"].group if outcome["school_class"] else "",
                    "debt": str(outstanding_for(enrollment)),
                })
            transaction.set_rollback(True)

        summary = {kind: sum(1 for row in rows if row["outcome"] == kind)
                   for kind in ("passe", "redouble", "diplômé", "indécis")}
        return Response({
            "ready": True,
            "next_year": target_year.name,
            "copied": copied,
            "summary": summary,
            "unassigned": sum(
                1 for row in rows
                if row["outcome"] in ("passe", "redouble") and not row["next_class"]
            ),
            "carried_debt": str(sum(Decimal(row["debt"]) for row in rows)),
            "students": rows,
        })


class AcademicSessionListView(SchoolScopedMixin, APIView):
    def get_academic_year(self, year_pk):
        try:
            return AcademicYear.objects.get(pk=year_pk, school=self.get_school())
        except AcademicYear.DoesNotExist:
            raise serializers.ValidationError({"academic_year": "Année académique introuvable."})

    def get(self, request, school_pk, year_pk):
        year = self.get_academic_year(year_pk)
        sessions = year.sessions.prefetch_related(
            "classes",
            Prefetch(
                "closure",
                queryset=SessionClosure.objects.select_related("closed_by"),
            ),
        )
        return Response(AcademicSessionSerializer(sessions, many=True).data)

    def post(self, request, school_pk, year_pk):
        self.ensure_school_life()
        year = self.get_academic_year(year_pk)
        if not year.is_active or year.is_closed:
            raise serializers.ValidationError({"academic_year": "Les sessions ne peuvent être créées que dans l’année académique active."})
        serializer = AcademicSessionSerializer(data={**request.data, "is_active": True}, context={"academic_year": year})
        serializer.is_valid(raise_exception=True)
        with transaction.atomic():
            serializer.save(academic_year=year)
        return Response(serializer.data, status=201)


class AcademicSessionDetailView(AcademicSessionListView):
    def get_object(self, year_pk, pk):
        year = self.get_academic_year(year_pk)
        try:
            return year.sessions.get(pk=pk), year
        except AcademicSession.DoesNotExist:
            raise serializers.ValidationError({"session": "Session académique introuvable."})

    def patch(self, request, school_pk, year_pk, pk):
        self.ensure_school_life()
        session, year = self.get_object(year_pk, pk)
        serializer = AcademicSessionSerializer(session, data=request.data, partial=True, context={"academic_year": year})
        serializer.is_valid(raise_exception=True)
        with transaction.atomic():
            serializer.save()
        return Response(serializer.data)

    def delete(self, request, school_pk, year_pk, pk):
        return Response({"detail": "Une session doit être clôturée et non supprimée."}, status=405)

    def post(self, request, school_pk, year_pk, pk):
        self.ensure_manager()
        session, _ = self.get_object(year_pk, pk)
        try:
            closure = close_session(session, user=request.user)
        except ClosureError as issue:
            return Response({"detail": str(issue)}, status=400)
        session.refresh_from_db()
        payload = AcademicSessionSerializer(session).data
        payload["closure"] = {
            "closed_at": closure.closed_at.isoformat(),
            "report_cards": closure.report_card_count,
            "grades": closure.grade_entry_count,
            "discipline": closure.discipline_count,
            "attendance_sessions": closure.attendance_session_count,
            "attendance_records": closure.attendance_record_count,
        }
        return Response(payload)


class StudentEnrollmentViewSet(SchoolScopedMixin, viewsets.ModelViewSet):
    serializer_class = StudentEnrollmentSerializer
    parser_classes = [JSONParser, MultiPartParser, FormParser]
    http_method_names = ["get", "post", "patch", "delete", "head", "options"]

    def get_academic_year(self):
        year_id = self.request.headers.get("X-Academic-Year-ID")
        if not year_id:
            raise serializers.ValidationError({"academic_year": "Sélectionnez une année académique."})
        try:
            academic_year = AcademicYear.objects.get(pk=year_id, school=self.get_school())
        except (AcademicYear.DoesNotExist, ValueError):
            raise serializers.ValidationError({"academic_year": "Année académique invalide."})
        return academic_year

    def get_queryset(self):
        return StudentEnrollment.objects.filter(
            school=self.get_school(),
            academic_year=self.get_academic_year(),
            status=StudentEnrollment.Status.ACTIVE,
        ).select_related("student", "guardian", "academic_year", "level", "school_class")

    def get_serializer_context(self):
        context = super().get_serializer_context()
        context.update(school=self.get_school(), academic_year=self.get_academic_year())
        return context

    def perform_create(self, serializer):
        self.ensure_school_life()
        academic_year = self.get_academic_year()
        if not academic_year.is_active or academic_year.is_closed:
            raise serializers.ValidationError({
                "academic_year": "Les inscriptions sont autorisées uniquement pour l’année active non clôturée."
            })
        serializer.save()

    def perform_destroy(self, instance):
        self.ensure_school_life()
        instance.status = StudentEnrollment.Status.CANCELLED
        instance.save(update_fields=["status"])

    def perform_update(self, serializer):
        self.ensure_school_life()
        serializer.save()

    def eligible_unassigned_queryset(self):
        return unassigned_students(self.get_school(), self.get_academic_year())

    @action(detail=False, methods=["get"], url_path="guardian-lookup")
    def guardian_lookup(self, request, school_pk=None):
        try:
            phone = normalize_togolese_phone(request.query_params.get("phone"))
        except serializers.ValidationError as exc:
            return Response({"phone": exc.detail}, status=400)
        guardian = CustomUser.objects.filter(role=CustomUser.Role.PARENT, phone=phone).first()
        if guardian is None:
            return Response({"found": False, "phone": phone})
        return Response({
            "found": True,
            "id": guardian.pk,
            "phone": guardian.phone,
            "last_name": guardian.last_name,
            "first_names": guardian.first_name,
            "profession": guardian.profession,
        })

    @action(detail=False, methods=["get"], url_path="unassigned")
    def unassigned(self, request, school_pk=None):
        queryset = self.eligible_unassigned_queryset().order_by(
            "level__order", "series", "student__last_name", "student__first_name",
        )
        # La moyenne manquante est retrouvée dans les bulletins de l'an
        # dernier : c'est elle qui classera l'élève, autant qu'elle s'affiche.
        students = attach_fallback_averages(
            self.get_school(), self.get_academic_year(), list(queryset),
        )
        return Response(self.get_serializer(students, many=True).data)

    @action(detail=False, methods=["post"], url_path="auto-assign")
    def auto_assign(self, request, school_pk=None):
        """Répartit les élèves sans classe selon les règles de l'établissement.

        Le calcul lui-même vit dans `assignment` : ici on rassemble les
        données — classes du niveau, effectifs, filles déjà inscrites — et on
        rend compte de ce qui a été fait.
        """
        self.ensure_school_life()
        school = self.get_school()
        academic_year = self.get_academic_year()
        if not academic_year.is_active or academic_year.is_closed:
            return Response({"academic_year": "L’année académique doit être active."}, status=400)

        students = list(self.eligible_unassigned_queryset())
        if not students:
            return Response({"message": "Aucun élève éligible sans classe.", "assigned": 0, "warnings": []})

        attach_fallback_averages(school, academic_year, students)
        rules = rules_for(school)
        classes = list(SchoolClass.objects.filter(
            school=school, academic_year=academic_year, is_active=True,
        ).annotate(
            effectif=Count(
                "student_enrollments",
                filter=Q(student_enrollments__status=StudentEnrollment.Status.ACTIVE),
                distinct=True,
            ),
        ).select_related("level"))

        # Filles déjà inscrites, classe par classe : l'équilibre des filles se
        # calcule sur le total à venir, pas sur les seules nouvelles venues.
        girls_before = dict(
            StudentEnrollment.objects
            .filter(
                school_class__in=classes,
                status=StudentEnrollment.Status.ACTIVE,
                student__gender=CustomUser.Gender.FEMALE,
            )
            .values("school_class")
            .annotate(total=Count("id"))
            .values_list("school_class", "total")
        )

        classes_by_group = {}
        for school_class in classes:
            classes_by_group.setdefault(
                group_key(school_class.level_id, school_class.series), [],
            ).append(school_class)

        students_by_group = {}
        for enrollment in students:
            students_by_group.setdefault(
                group_key(enrollment.level_id, enrollment.series), [],
            ).append(enrollment)

        warnings, assignments, summary = [], [], []
        for key, grouped_students in students_by_group.items():
            grouped_classes = classes_by_group.get(key, [])
            level_name = grouped_students[0].level.name if grouped_students[0].level else "Niveau inconnu"
            group_label = f"{level_name} — {grouped_students[0].series or 'sans série'}"
            if not grouped_classes:
                warnings.append(
                    f"{group_label} : aucune classe active correspondante pour "
                    f"{len(grouped_students)} élève(s)."
                )
                continue

            allocations = distribute_group(
                grouped_classes, grouped_students, girls_before, rules,
            )
            ordered_classes = sorted(grouped_classes, key=natural_class_key)
            assigned_in_group = 0
            for school_class in ordered_classes:
                for enrollment in allocations[school_class.pk]:
                    enrollment.school_class = school_class
                    assignments.append(enrollment)
                    assigned_in_group += 1

            summary.append({
                "group": group_label,
                "assigned": assigned_in_group,
                "classes": [
                    {
                        "id": school_class.pk,
                        "name": school_class.name,
                        "added": len(allocations[school_class.pk]),
                        "total": school_class.effectif + len(allocations[school_class.pk]),
                        "capacity": school_class.maximum_capacity,
                        "girls": (
                            girls_before.get(school_class.pk, 0)
                            + sum(1 for row in allocations[school_class.pk] if is_girl(row))
                        ),
                        # Le maximum a-t-il dû être dépassé pour placer tout le
                        # monde ? L'utilisateur doit le voir sans le déduire.
                        "over_capacity": (
                            school_class.effectif + len(allocations[school_class.pk])
                            > school_class.maximum_capacity
                        ),
                    }
                    for school_class in ordered_classes
                ],
            })
            if assigned_in_group < len(grouped_students):
                warnings.append(
                    f"{group_label} : capacité insuffisante, "
                    f"{len(grouped_students) - assigned_in_group} élève(s) restent sans classe."
                )

        with transaction.atomic():
            StudentEnrollment.objects.bulk_update(assignments, ["school_class"])

        return Response({
            "message": f"{len(assignments)} élève(s) réparti(s) automatiquement.",
            "assigned": len(assignments),
            "remaining": len(students) - len(assignments),
            "warnings": warnings,
            "summary": summary,
        })

    @action(detail=False, methods=["post"], url_path="import", parser_classes=[MultiPartParser, FormParser])
    def import_file(self, request, school_pk=None):
        self.ensure_school_life()
        school = self.get_school()
        academic_year = self.get_academic_year()
        if not academic_year.is_active or academic_year.is_closed:
            return Response({"academic_year": "L’année académique doit être active."}, status=400)
        uploaded_file = request.FILES.get("file")
        if not uploaded_file:
            return Response({"file": "Sélectionnez un fichier Excel ou CSV."}, status=400)

        try:
            if uploaded_file.name.lower().endswith(".csv"):
                content = uploaded_file.read().decode("utf-8-sig")
                rows = list(csv.DictReader(io.StringIO(content), delimiter=";" if ";" in content.splitlines()[0] else ","))
            elif uploaded_file.name.lower().endswith(".xlsx"):
                try:
                    from openpyxl import load_workbook
                except ImportError:
                    rows = read_xlsx_without_dependency(uploaded_file)
                else:
                    sheet = load_workbook(uploaded_file, read_only=True, data_only=True).active
                    values = list(sheet.iter_rows(values_only=True))
                    if not values:
                        rows = []
                    else:
                        headers = [str(value or "").strip() for value in values[0]]
                        rows = [dict(zip(headers, values_row)) for values_row in values[1:] if any(value not in (None, "") for value in values_row)]
            else:
                return Response({"file": "Formats acceptés : .xlsx et .csv."}, status=400)
        except (UnicodeDecodeError, ValueError, OSError, zipfile.BadZipFile, ET.ParseError, KeyError) as exc:
            return Response({"file": f"Impossible de lire le fichier : {exc}"}, status=400)

        aliases = {
            "matricule": "enrollment_number", "numero_matricule": "enrollment_number",
            "nom": "last_name", "prenoms": "first_names", "prénoms": "first_names",
            "genre": "gender", "date_naissance": "date_of_birth", "date_de_naissance": "date_of_birth",
            "cycle": "cycle", "niveau": "level", "serie": "series", "série": "series", "statut": "student_status", "status": "student_status", "classe": "school_class", "nom_classe": "school_class", "nom_de_la_classe": "school_class",
            "allergie": "health_information", "allergies": "health_information",
            "souci_de_sante": "health_information", "soucis_de_sante": "health_information", "probleme_de_sante": "health_information",
            "informations_de_sante": "health_information", "allergies_et_soucis_de_sante": "health_information",
            "moyenne": "previous_average", "moyenne_precedente": "previous_average",
            "moyenne_annee_ecoulee": "previous_average", "moyenne_scolaire": "previous_average",
            "moyenne_scolaire_de_l_annee_ecoulee": "previous_average",
            "telephone_tuteur": "guardian_phone", "numero_tuteur": "guardian_phone",
            "nom_tuteur": "guardian_last_name", "prenom_tuteur": "guardian_first_names",
            "prenoms_tuteur": "guardian_first_names", "profession_tuteur": "guardian_profession",
        }
        normalized_rows = []
        errors = []
        seen_numbers = set()
        next_sequence = StudentEnrollment.objects.filter(school=school).count() + 1

        for index, source_row in enumerate(rows, start=2):
            row = {aliases.get(slugify(str(key)).replace("-", "_"), slugify(str(key)).replace("-", "_")): value for key, value in source_row.items()}
            level_name = str(row.get("level") or "").strip()
            level = SchoolLevel.objects.filter(school=school, name__iexact=level_name, is_active=True).first()
            cycle_value = slugify(str(row.get("cycle") or "").strip())
            cycle_map = {"primaire": SchoolLevel.Stage.PRIMARY, "college": SchoolLevel.Stage.MIDDLE, "lycee": SchoolLevel.Stage.HIGH}
            selected_cycle = cycle_map.get(cycle_value)
            class_name = str(row.get("school_class") or "").strip()
            school_class = SchoolClass.objects.filter(
                school=school, academic_year=academic_year, group__iexact=class_name, is_active=True,
            ).first() if class_name else None
            series = str(row.get("series") or "").strip()
            student_status = slugify(str(row.get("student_status") or "nouveau").strip()) or "nouveau"
            if not series and school_class:
                series = school_class.series
            number = str(row.get("enrollment_number") or "").strip().upper().replace(" ", "")
            if not number:
                while True:
                    number = f"{academic_year.start_date.year}-{next_sequence:04d}"
                    next_sequence += 1
                    if number not in seen_numbers and not StudentEnrollment.objects.filter(school=school, enrollment_number__iexact=number).exists():
                        break
            gender_value = str(row.get("gender") or "").strip().lower()
            gender = {"m": "M", "masculin": "M", "f": "F", "feminin": "F", "féminin": "F"}.get(gender_value, "")
            birth_date = row.get("date_of_birth")
            if isinstance(birth_date, (datetime, date)):
                birth_date = birth_date.strftime("%Y-%m-%d")
            payload = {
                "enrollment_number": number,
                "last_name": str(row.get("last_name") or "").strip(),
                "first_names": str(row.get("first_names") or "").strip(),
                "gender": gender,
                "date_of_birth": str(birth_date or "").strip(),
                "health_information": str(row.get("health_information") or "").strip(),
                "level": level.pk if level else None,
                "series": series,
                "student_status": student_status,
                "previous_average": str(row.get("previous_average") or "").strip().replace(",", ".") or None,
                "guardian_phone": str(row.get("guardian_phone") or "").strip(),
                "guardian_last_name": str(row.get("guardian_last_name") or "").strip(),
                "guardian_first_names": str(row.get("guardian_first_names") or "").strip(),
                "guardian_profession": str(row.get("guardian_profession") or "").strip(),
                "school_class": school_class.pk if school_class else None,
            }
            import_rule_errors = {}
            if not selected_cycle:
                import_rule_errors["cycle"] = ["Cycle obligatoire ou invalide. Choix : Primaire, Collège ou Lycée."]
            elif level and level.stage != selected_cycle:
                import_rule_errors["cycle"] = [f"Le niveau « {level_name} » appartient au cycle {level.get_stage_display()}, pas au cycle indiqué."]
            if selected_cycle in (SchoolLevel.Stage.PRIMARY, SchoolLevel.Stage.MIDDLE) and series:
                import_rule_errors["serie"] = [f"Le cycle {'Primaire' if selected_cycle == SchoolLevel.Stage.PRIMARY else 'Collège'} ne doit pas avoir de série."]
            if number.casefold() in seen_numbers:
                errors.append({"ligne": index, "erreurs": {"matricule": ["Matricule répété dans le fichier."]}})
                continue
            seen_numbers.add(number.casefold())
            serializer = self.get_serializer(data=payload)
            if serializer.is_valid():
                row_errors = dict(import_rule_errors)
                if class_name and not school_class:
                    row_errors["classe"] = [f"La classe « {class_name} » n’existe pas dans cette école pour l’année {academic_year.name}. Vérifiez son nom exact."]
                elif school_class and level and school_class.level_id != level.id:
                    row_errors["classe"] = [f"La classe « {class_name} » appartient au niveau {school_class.level.name}, pas au niveau {level_name}."]
                if row_errors:
                    errors.append({"ligne": index, "erreurs": row_errors})
                else:
                    normalized_rows.append(serializer)
            else:
                row_errors = {**dict(serializer.errors), **import_rule_errors}
                if not level:
                    row_errors["niveau"] = [f"Niveau « {level_name} » introuvable dans cette école."]
                if class_name and not school_class:
                    row_errors["classe"] = [f"La classe « {class_name} » n’existe pas dans cette école pour l’année {academic_year.name}. Vérifiez son nom exact."]
                elif school_class and level and school_class.level_id != level.id:
                    row_errors["classe"] = [f"La classe « {class_name} » appartient au niveau {school_class.level.name}, pas au niveau {level_name}."]
                errors.append({"ligne": index, "erreurs": row_errors})

        if not rows:
            return Response({"file": "Le fichier ne contient aucun élève."}, status=400)
        if errors:
            return Response({
                "message": "Import annulé : aucune donnée du fichier n’a été enregistrée. Corrigez les lignes indiquées puis réessayez.",
                "errors": errors,
            }, status=400)

        with transaction.atomic():
            enrollments = [serializer.save() for serializer in normalized_rows]
        return Response({
            "message": f"{len(enrollments)} élève(s) importé(s) avec succès.",
            "count": len(enrollments),
            "enrollments": self.get_serializer(enrollments, many=True).data,
        }, status=201)


class AssignmentSettingsView(SchoolScopedMixin, APIView):
    """Règles de répartition des élèves sans classe.

    Lisibles par qui répartit — le secrétariat en fait partie — mais
    modifiables par la seule direction : décider que les meilleurs vont devant
    ou que les effectifs priment est une politique d'établissement.
    """

    BOOLEANS = [
        "youngest_first", "best_first", "balance_headcount",
        "balance_girls", "allow_overflow",
    ]
    NUMBERS = ["reserved_excellent", "overflow_margin"]
    # Bornes basses des tranches de moyennes, seules clés acceptées en limite.
    BAND_KEYS = {band[1] for band in AVERAGE_BANDS}

    def get_academic_year(self):
        try:
            return AcademicYear.objects.get(
                pk=self.request.headers.get("X-Academic-Year-ID"), school=self.get_school(),
            )
        except (AcademicYear.DoesNotExist, ValueError, TypeError):
            raise serializers.ValidationError(
                {"academic_year": "Sélectionnez une année académique valide."}
            )

    def clean_band_limits(self, submitted):
        """Vérifie et range les plafonds : un cursus de l'école, une tranche connue.

        La clé est « niveau|série » — « 12|A4 », ou « 9| » pour un niveau sans
        série. Les zéros ne sont pas conservés, ni les cursus qui n'en portent
        aucun : le réglage enregistré ne garde que ce qui agit.
        """
        if not isinstance(submitted, dict):
            raise serializers.ValidationError(
                {"band_limits": "Attendu : un plafond par cursus et par tranche."}
            )
        known = set(
            SchoolLevel.objects
            .filter(school=self.get_school())
            .values_list("id", flat=True)
        )
        cleaned = {}
        for stream, bands in submitted.items():
            level_key, _, series = str(stream).partition("|")
            try:
                level_id = int(level_key)
            except (TypeError, ValueError):
                raise serializers.ValidationError(
                    {"band_limits": f"Niveau inconnu : {level_key}."}
                )
            if level_id not in known:
                raise serializers.ValidationError(
                    {"band_limits": f"Niveau inconnu : {level_key}."}
                )
            if len(series) > 20:
                raise serializers.ValidationError(
                    {"band_limits": f"Série invalide : {series}."}
                )
            if not isinstance(bands, dict):
                raise serializers.ValidationError(
                    {"band_limits": "Attendu : un plafond par tranche de moyenne."}
                )
            for band, value in bands.items():
                try:
                    if Decimal(str(band)) not in self.BAND_KEYS:
                        raise InvalidOperation
                except (InvalidOperation, TypeError, ValueError):
                    raise serializers.ValidationError(
                        {"band_limits": f"Tranche inconnue : {band}."}
                    )
                try:
                    limit = int(value)
                except (TypeError, ValueError):
                    raise serializers.ValidationError(
                        {"band_limits": "Saisissez un nombre entier par tranche."}
                    )
                if not 0 <= limit <= 200:
                    raise serializers.ValidationError(
                        {"band_limits": "Saisissez un nombre entre 0 et 200."}
                    )
                if limit:
                    key = stream_key(level_id, series)
                    cleaned.setdefault(key, {})[str(Decimal(str(band)))] = limit
        return cleaned

    def distribution(self):
        """Combien d'élèves à répartir dans chaque tranche de moyenne.

        Sert à choisir le seuil de la réserve en connaissance de cause : mettre
        la barre à 16 n'a pas le même effet selon qu'il y a trois élèves
        au-dessus ou soixante.
        """
        school = self.get_school()
        year = self.get_academic_year()
        students = attach_fallback_averages(
            school, year, list(unassigned_students(school, year)),
        )
        return band_counts(students)

    def serialize(self, rules, distribution=None):
        counts = distribution or self.distribution()
        return {
            **{field: getattr(rules, field) for field in self.BOOLEANS},
            **{field: getattr(rules, field) for field in self.NUMBERS},
            "excellent_minimum": str(rules.excellent_minimum),
            "can_configure": self.has_role_for_school(self.get_school(), DIRECTION_ROLES),
            **counts,
            # Plafonds tels qu'ils s'affichent : une case par cursus et par
            # tranche, zéro là où rien n'est limité. Les réglages enregistrés
            # pour un cursus absent de la liste sont rendus tels quels, sinon
            # le prochain enregistrement les effacerait.
            "band_limits": {
                **{
                    key: {str(band): value for band, value in bands.items()}
                    for key, bands in (rules.band_limits or {}).items()
                    if isinstance(bands, dict)
                },
                **{
                    row["key"]: {
                        str(key): band_limits(rules, row["level"], row["series"]).get(str(key), 0)
                        for key in self.BAND_KEYS
                    }
                    for row in counts["levels"] if row["key"]
                },
            },
        }

    def get(self, request, school_pk):
        self.ensure_school_life()
        return Response(self.serialize(rules_for(self.get_school())))

    def put(self, request, school_pk):
        self.ensure_manager()
        rules = rules_for(self.get_school())
        for field in self.BOOLEANS:
            if field in request.data:
                setattr(rules, field, bool(request.data[field]))
        for field in self.NUMBERS:
            if field in request.data:
                try:
                    value = int(request.data[field])
                except (TypeError, ValueError):
                    raise serializers.ValidationError({field: "Saisissez un nombre entier."})
                if not 0 <= value <= 50:
                    raise serializers.ValidationError({field: "Saisissez un nombre entre 0 et 50."})
                setattr(rules, field, value)
        if "excellent_minimum" in request.data:
            try:
                minimum = Decimal(str(request.data["excellent_minimum"]))
            except (InvalidOperation, TypeError):
                raise serializers.ValidationError({"excellent_minimum": "Saisissez une moyenne."})
            if not 0 <= minimum <= 20:
                raise serializers.ValidationError(
                    {"excellent_minimum": "La moyenne doit être comprise entre 0 et 20."}
                )
            rules.excellent_minimum = minimum
        if "band_limits" in request.data:
            rules.band_limits = self.clean_band_limits(request.data["band_limits"])
        rules.save()
        return Response(self.serialize(rules))


class StudentJourneyView(SchoolScopedMixin, APIView):
    """Recherche d'un élève et dossier de son parcours dans l'établissement.

    Sans `student`, la vue cherche par nom ou matricule ; avec, elle rend le
    dossier complet. La recherche ne se limite pas à l'année active : c'est
    justement pour retrouver un ancien élève qu'on ouvre un parcours.
    """

    def get(self, request, school_pk):
        school = self.get_school()
        student_id = request.query_params.get("student")
        if not student_id:
            return Response({
                "results": search_students(school, request.query_params.get("q", "")),
            })
        try:
            if not StudentEnrollment.objects.filter(
                school=school, student_id=student_id,
            ).exists():
                raise CustomUser.DoesNotExist
            return Response(student_journey(school, int(student_id)))
        except (CustomUser.DoesNotExist, TypeError, ValueError):
            raise serializers.ValidationError(
                {"student": "Aucun élève de cet établissement ne porte cet identifiant."}
            )


class EnrollmentNumberSuggestionView(SchoolScopedMixin, APIView):
    def get(self, request, school_pk):
        school = self.get_school()
        year_id = request.headers.get("X-Academic-Year-ID")
        try:
            academic_year = AcademicYear.objects.get(pk=year_id, school=school)
        except (AcademicYear.DoesNotExist, ValueError, TypeError):
            return Response({"detail": "Sélectionnez une année académique valide."}, status=400)
        sequence = StudentEnrollment.objects.filter(school=school).count() + 1
        while True:
            number = f"{academic_year.start_date.year}-{sequence:04d}"
            if not StudentEnrollment.objects.filter(
                school=school, enrollment_number__iexact=number
            ).exists():
                return Response({"enrollment_number": number})
            sequence += 1


class SchoolLevelViewSet(SchoolScopedMixin, viewsets.ReadOnlyModelViewSet):
    serializer_class = SchoolLevelSerializer

    def get_queryset(self):
        return SchoolLevel.objects.filter(school=self.get_school(), is_active=True)


class SchoolClassViewSet(SchoolScopedMixin, viewsets.ModelViewSet):
    serializer_class = SchoolClassSerializer

    def get_academic_year(self):
        year_id = self.request.headers.get("X-Academic-Year-ID")
        try:
            return AcademicYear.objects.get(pk=year_id, school=self.get_school())
        except (AcademicYear.DoesNotExist, ValueError, TypeError):
            raise serializers.ValidationError({"academic_year": "Sélectionnez une année académique valide."})

    def get_queryset(self):
        queryset = SchoolClass.objects.filter(school=self.get_school(), academic_year=self.get_academic_year())
        if self.request.query_params.get("include_inactive", "").lower() not in {"1", "true", "yes"}:
            queryset = queryset.filter(is_active=True)
        return queryset.annotate(
            effectif=Count("student_enrollments", filter=Q(student_enrollments__status=StudentEnrollment.Status.ACTIVE), distinct=True),
        ).select_related("level", "academic_year", "homeroom_teacher").prefetch_related("subject_configurations__subject")

    def get_serializer_context(self):
        context = super().get_serializer_context()
        context.update(school=self.get_school(), academic_year=self.get_academic_year())
        return context

    def perform_create(self, serializer):
        self.ensure_school_life()
        year = self.get_academic_year()
        if not year.is_active or year.is_closed:
            raise serializers.ValidationError({"academic_year": "Les classes peuvent être créées uniquement dans l’année active."})
        serializer.save(school=self.get_school(), academic_year=year)

    def perform_update(self, serializer):
        self.ensure_school_life()
        serializer.save()

    def perform_destroy(self, instance):
        self.ensure_school_life()
        instance.is_active = False
        instance.save(update_fields=["is_active"])


class IsSuperUser(BasePermission):
    def has_permission(self, request, view):
        return bool(request.user and request.user.is_authenticated and request.user.is_superuser)


class OwnerViewSet(viewsets.ModelViewSet):
    serializer_class = CustomUserSerializer
    permission_classes = [IsSuperUser]
    http_method_names = ["get", "post", "patch", "delete", "head", "options"]

    def get_queryset(self):
        return CustomUser.objects.filter(role=CustomUser.Role.OWNER, is_archived=False)

    def perform_create(self, serializer):
        serializer.save(role=CustomUser.Role.OWNER)

    def perform_destroy(self, instance):
        instance.is_archived = True
        instance.is_active = False
        instance.save(update_fields=["is_archived", "is_active"])


class ParentListView(SchoolScopedMixin, APIView):
    """Parents et tuteurs de l'école, avec les élèves dont ils ont la charge."""

    def get(self, request, school_pk):
        school = self.get_school()
        enrollments = StudentEnrollment.objects.filter(
            school=school, guardian__isnull=False, status=StudentEnrollment.Status.ACTIVE,
        ).select_related("guardian", "student", "school_class", "school_class__level", "academic_year")

        year_id = request.headers.get("X-Academic-Year-ID")
        if year_id:
            enrollments = enrollments.filter(academic_year_id=year_id)

        parents = {}
        for enrollment in enrollments:
            guardian = enrollment.guardian
            entry = parents.setdefault(guardian.id, {
                "id": guardian.id,
                "username": guardian.username,
                "last_name": guardian.last_name,
                "first_names": guardian.first_name,
                "phone": guardian.phone,
                "email": guardian.email,
                "profession": guardian.profession,
                "address": guardian.address,
                "children": [],
            })
            entry["children"].append({
                "enrollment_id": enrollment.id,
                "student_id": enrollment.student_id,
                "name": enrollment.student.get_full_name(),
                "matricule": enrollment.enrollment_number,
                "class_name": enrollment.school_class.group if enrollment.school_class else None,
                "level": enrollment.school_class.level.name if enrollment.school_class else None,
            })

        rows = sorted(parents.values(), key=lambda item: (item["last_name"] or "", item["first_names"] or ""))
        students_total = StudentEnrollment.objects.filter(
            school=school, status=StudentEnrollment.Status.ACTIVE,
            **({"academic_year_id": year_id} if year_id else {}),
        ).count()
        return Response({
            "parents": rows,
            "students_total": students_total,
            "students_with_guardian": sum(len(item["children"]) for item in rows),
        })


class TimetableMixin(SchoolScopedMixin):
    def get_academic_year(self):
        try:
            return AcademicYear.objects.get(
                pk=self.request.headers.get("X-Academic-Year-ID"), school=self.get_school(),
            )
        except (AcademicYear.DoesNotExist, ValueError, TypeError):
            raise serializers.ValidationError({"academic_year": "Sélectionnez une année académique valide."})

    def can_manage(self):
        return self.has_role_for_school(self.get_school(), DIRECTION_ROLES)

    def ensure_manager_access(self):
        if not self.can_manage():
            raise serializers.ValidationError(
                {"permission": "Seule la direction de l’établissement peut gérer l’emploi du temps."}
            )

    def serialize(self, timetable):
        slots = timetable.slots.select_related(
            "school_class", "school_class__level", "class_subject__subject", "teacher",
        ).all()
        return {
            "id": timetable.id,
            "status": timetable.status,
            "status_label": timetable.get_status_display(),
            "is_validated": timetable.is_validated,
            "days_per_week": timetable.days_per_week,
            "period_duration": timetable.period_duration,
            "generated_at": timetable.generated_at,
            "validated_at": timetable.validated_at,
            "options": {
                "enforce_paired_hours": timetable.enforce_paired_hours,
                "enforce_single_hour_middle": timetable.enforce_single_hour_middle,
                "enforce_day_spacing": timetable.enforce_day_spacing,
                "enforce_max_two_hours": timetable.enforce_max_two_hours,
                "skip_primary": timetable.skip_primary,
            },
            "days_without_afternoon": timetable.days_without_afternoon or [],
            # Volume horaire de la classe la plus chargée : sert à prévenir
            # l'utilisateur quand la grille devient trop courte.
            "max_class_hours": ClassSubject.objects.filter(
                school_class__academic_year=timetable.academic_year_id,
            ).values("school_class").annotate(total=Sum("weekly_hours")).aggregate(
                peak=Max("total"),
            )["peak"] or 0,
            "periods": [{
                "id": period.id,
                "label": period.label,
                "kind": period.kind,
                "kind_label": period.get_kind_display(),
                "start_time": period.start_time,
                "end_time": period.end_time,
                "order": period.order,
            } for period in timetable.periods.all()],
            "restrictions": [{
                "id": restriction.id,
                "subject": restriction.subject_id,
                "subject_name": restriction.subject.name,
                "period": restriction.period_id,
                "day": restriction.day,
            } for restriction in timetable.restrictions.select_related("subject").all()],
            "excluded_classes": [{
                "id": exclusion.id,
                "subject": exclusion.subject_id,
                "subject_name": exclusion.subject.name,
                "classes": [item.id for item in exclusion.classes.all()],
                "class_names": [item.group for item in exclusion.classes.all()],
            } for exclusion in timetable.excluded_classes.select_related("subject").prefetch_related("classes")],
            "class_groups": [{
                "id": group.id,
                "subject": group.subject_id,
                "subject_name": group.subject.name,
                "classes": [item.id for item in group.classes.all()],
                "class_names": [item.group for item in group.classes.all()],
            } for group in timetable.class_groups.select_related("subject").prefetch_related("classes")],
            "slots": [{
                "id": slot.id,
                "class_id": slot.school_class_id,
                "class_name": slot.school_class.group,
                "level": slot.school_class.level.name,
                "subject": slot.class_subject.subject.name,
                "subject_code": slot.class_subject.subject.code,
                "teacher": slot.teacher.get_full_name() if slot.teacher else None,
                "teacher_id": slot.teacher_id,
                "day": slot.day,
                "day_label": DAY_LABELS.get(slot.day, ""),
                "start_time": slot.start_time,
                "end_time": slot.end_time,
            } for slot in slots],
        }


class TimetableView(TimetableMixin, APIView):
    """Consultation, génération et suppression de l'emploi du temps annuel."""

    def get(self, request, school_pk):
        year = self.get_academic_year()
        timetable = Timetable.objects.filter(school=self.get_school(), academic_year=year).first()
        if not timetable:
            return Response({"timetable": None, "can_manage": self.can_manage()})
        return Response({"timetable": self.serialize(timetable), "can_manage": self.can_manage()})

    def post(self, request, school_pk):
        """Génère (ou régénère) l'emploi du temps de l'année."""
        self.ensure_manager_access()
        school = self.get_school()
        year = self.get_academic_year()

        timetable = Timetable.objects.filter(school=school, academic_year=year).first()
        if timetable and timetable.is_validated:
            raise serializers.ValidationError(
                {"timetable": "L’emploi du temps de cette année est validé. Supprimez-le pour en générer un nouveau."}
            )
        if not ClassSubject.objects.filter(school_class__academic_year=year).exists():
            raise serializers.ValidationError(
                {"timetable": "Aucune matière n’est configurée pour cette année académique."}
            )

        if not timetable:
            timetable = Timetable.objects.create(school=school, academic_year=year, generated_by=request.user)
        for field in ("days_per_week", "period_duration"):
            if field in request.data:
                setattr(timetable, field, int(request.data[field]))
        timetable.save()

        _, unplaced, relaxed, attempts, floor = regenerate(timetable)
        _, overflow = irreducible_deficit(timetable)
        timetable.refresh_from_db()
        return Response({
            "timetable": self.serialize(timetable),
            "unplaced": describe_unplaced(unplaced),
            "relaxed": relaxed,
            "attempts": attempts,
            "irreducible": floor,
            "overloaded_classes": overflow,
            "can_manage": True,
        })

    def delete(self, request, school_pk):
        """Vide la grille sans toucher au paramétrage.

        Seules les cases générées sont effacées : créneaux horaires,
        interdictions, exclusions, regroupements et options restent tels que
        l'utilisateur les a réglés. Supprimer la ligne `Timetable` les
        emporterait tous en cascade.
        """
        self.ensure_manager_access()
        year = self.get_academic_year()
        timetable = Timetable.objects.filter(school=self.get_school(), academic_year=year).first()
        if not timetable:
            raise serializers.ValidationError({"timetable": "Aucun emploi du temps à supprimer."})
        if timetable.is_validated:
            raise serializers.ValidationError(
                {"timetable": "Cet emploi du temps est validé. Repassez-le en brouillon pour le vider."}
            )
        timetable.slots.all().delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


class TimetableSetupView(TimetableMixin, APIView):
    """Paramétrage préalable : créneaux horaires, contraintes actives et interdictions."""

    OPTION_FIELDS = [
        "enforce_paired_hours", "enforce_single_hour_middle",
        "enforce_day_spacing", "enforce_max_two_hours", "skip_primary",
    ]

    def get_or_create_timetable(self):
        school = self.get_school()
        year = self.get_academic_year()
        timetable = Timetable.objects.filter(school=school, academic_year=year).first()
        if not timetable:
            timetable = Timetable.objects.create(
                school=school, academic_year=year, generated_by=self.request.user,
            )
        if not timetable.periods.exists():
            create_default_periods(timetable)
        return timetable

    def get(self, request, school_pk):
        timetable = self.get_or_create_timetable()
        year = self.get_academic_year()
        subjects = Subject.objects.filter(school=self.get_school(), is_active=True).order_by("name")
        classes = SchoolClass.objects.filter(academic_year=year).select_related("level").order_by("group")

        # Enseignant de chaque couple (classe, matière) : l'interface s'en sert
        # pour n'autoriser le regroupement qu'entre classes au même professeur.
        teachers = {}
        for link in TeacherAssignmentSubject.objects.filter(
            assignment__academic_year=year,
        ).select_related("assignment__teacher", "class_subject"):
            key = f"{link.class_subject.school_class_id}-{link.class_subject.subject_id}"
            teachers[key] = {
                "id": link.assignment.teacher_id,
                "name": link.assignment.teacher.get_full_name() or link.assignment.teacher.last_name,
            }

        return Response({
            "timetable": self.serialize(timetable),
            "subjects": [{"id": subject.id, "name": subject.name, "code": subject.code} for subject in subjects],
            "classes": [{
                "id": school_class.id,
                "name": school_class.group,
                "level": school_class.level.name,
            } for school_class in classes],
            "class_subject_teachers": teachers,
            "can_manage": self.can_manage(),
        })

    @transaction.atomic
    def put(self, request, school_pk):
        self.ensure_manager_access()
        timetable = self.get_or_create_timetable()
        if timetable.is_validated:
            raise serializers.ValidationError(
                {"timetable": "Cet emploi du temps est validé. Supprimez-le pour modifier le paramétrage."}
            )

        for field in self.OPTION_FIELDS:
            if field in request.data:
                setattr(timetable, field, bool(request.data[field]))
        if "days_per_week" in request.data:
            days = int(request.data["days_per_week"])
            if not 1 <= days <= 7:
                raise serializers.ValidationError({"days_per_week": "Indiquez entre 1 et 7 jours."})
            timetable.days_per_week = days

        if "days_without_afternoon" in request.data:
            values = request.data["days_without_afternoon"] or []
            if not isinstance(values, list):
                raise serializers.ValidationError(
                    {"days_without_afternoon": "Indiquez une liste de jours."}
                )
            closed = set()
            for value in values:
                try:
                    day = int(value)
                except (TypeError, ValueError):
                    raise serializers.ValidationError(
                        {"days_without_afternoon": "Jour invalide."}
                    )
                if not 0 <= day <= 6:
                    raise serializers.ValidationError(
                        {"days_without_afternoon": "Les jours vont de 0 (lundi) à 6 (dimanche)."}
                    )
                closed.add(day)
            timetable.days_without_afternoon = sorted(closed)

        timetable.save()

        periods = request.data.get("periods")
        # Correspondance entre l'identifiant envoyé par le client et le créneau
        # réellement enregistré : les interdictions s'y rattachent ensuite.
        period_aliases = {}
        if periods is not None:
            if not periods:
                raise serializers.ValidationError({"periods": "Définissez au moins un créneau."})
            existing = {period.id: period for period in timetable.periods.all()}
            rows = []
            for order, item in enumerate(periods, start=1):
                start = str(item.get("start_time", ""))[:5]
                end = str(item.get("end_time", ""))[:5]
                if not start or not end or start >= end:
                    raise serializers.ValidationError(
                        {"periods": f"Créneau {order} : l’heure de fin doit suivre l’heure de début."}
                    )
                kind = item.get("kind")
                if kind not in dict(TimetablePeriod.Kind.choices):
                    kind = TimetablePeriod.Kind.COURSE
                rows.append({
                    "id": item.get("id"),
                    "label": str(item.get("label", ""))[:60],
                    "kind": kind,
                    "start": start,
                    "end": end,
                    "order": order,
                })

            starts = [row["start"] for row in rows]
            if len(set(starts)) != len(starts):
                raise serializers.ValidationError(
                    {"periods": "Deux créneaux ne peuvent pas commencer à la même heure."}
                )

            # Supprimer d'abord les créneaux abandonnés : sans cela leurs horaires
            # entrent en collision avec ceux qu'on s'apprête à écrire.
            reused = {row["id"] for row in rows if row["id"] in existing}
            timetable.periods.exclude(pk__in=reused).delete()

            # Les lignes réutilisées sont d'abord écartées sur des horaires et des
            # ordres libres : sans cela, permuter deux créneaux violerait les
            # contraintes d'unicité pendant la réécriture.
            for offset, period in enumerate(existing[pk] for pk in reused):
                period.start_time = time(offset // 60, offset % 60, 0)
                period.end_time = time(offset // 60, offset % 60, 30)
                period.order = 100 + offset
                period.save(update_fields=["start_time", "end_time", "order"])

            for row in rows:
                # Réutiliser la ligne existante préserve les interdictions qui la visent.
                period = existing[row["id"]] if row["id"] in reused else TimetablePeriod(timetable=timetable)
                period.label, period.kind = row["label"], row["kind"]
                period.start_time, period.end_time = row["start"], row["end"]
                period.order = row["order"]
                period.save()
                if row["id"]:
                    period_aliases[row["id"]] = period.id

        restrictions = request.data.get("restrictions")
        if restrictions is not None:
            valid_periods = set(timetable.periods.values_list("id", flat=True))
            valid_subjects = set(
                Subject.objects.filter(school=self.get_school()).values_list("id", flat=True)
            )
            rows = []
            seen = set()
            for item in restrictions:
                subject_id = item.get("subject")
                period_id = period_aliases.get(item.get("period"), item.get("period"))
                day = item.get("day")
                if subject_id not in valid_subjects or period_id not in valid_periods:
                    continue
                key = (subject_id, period_id, day)
                if key in seen:
                    continue
                seen.add(key)
                rows.append(SubjectPeriodRestriction(
                    timetable=timetable, subject_id=subject_id, period_id=period_id,
                    day=int(day) if day is not None and day != "" else None,
                ))
            timetable.restrictions.all().delete()
            SubjectPeriodRestriction.objects.bulk_create(rows)

        year = self.get_academic_year()
        valid_classes = {
            item.id: item
            for item in SchoolClass.objects.filter(academic_year=year)
        }

        valid_subjects = set(
            Subject.objects.filter(school=self.get_school()).values_list("id", flat=True)
        )

        excluded_classes = request.data.get("excluded_classes")
        if excluded_classes is not None:
            prepared_exclusions = []
            for item in excluded_classes:
                subject_id = item.get("subject")
                members = [
                    class_id for class_id in dict.fromkeys(item.get("classes") or [])
                    if class_id in valid_classes
                ]
                if subject_id not in valid_subjects or not members:
                    continue
                prepared_exclusions.append((subject_id, members))

            timetable.excluded_classes.all().delete()
            for subject_id, members in prepared_exclusions:
                exclusion = ExcludedTimetableClass.objects.create(
                    timetable=timetable, subject_id=subject_id,
                )
                exclusion.classes.set(members)

        class_groups = request.data.get("class_groups")
        if class_groups is not None:
            prepared = []
            for item in class_groups:
                subject_id = item.get("subject")
                members = [
                    class_id for class_id in dict.fromkeys(item.get("classes") or [])
                    if class_id in valid_classes
                ]
                if subject_id not in valid_subjects or len(members) < 2:
                    continue

                # Un regroupement n'a de sens que si un seul enseignant assure
                # la matière dans toutes les classes réunies.
                teachers = set()
                for class_id in members:
                    link = TeacherAssignmentSubject.objects.filter(
                        class_subject__school_class_id=class_id,
                        class_subject__subject_id=subject_id,
                        assignment__academic_year=year,
                    ).select_related("assignment").first()
                    teachers.add(link.assignment.teacher_id if link else None)
                if len(teachers) > 1:
                    names = ", ".join(valid_classes[class_id].group for class_id in members)
                    raise serializers.ValidationError({"class_groups": (
                        f"Les classes {names} n’ont pas le même enseignant pour cette matière : "
                        "elles ne peuvent pas être réunies."
                    )})
                prepared.append((subject_id, members))

            timetable.class_groups.all().delete()
            for subject_id, members in prepared:
                group = ClassGroupSession.objects.create(timetable=timetable, subject_id=subject_id)
                group.classes.set(members)

        timetable.refresh_from_db()
        return Response(self.serialize(timetable))


class TimetableValidationView(TimetableMixin, APIView):
    """Verrouille l'emploi du temps : une fois validé il n'est plus régénérable."""

    def post(self, request, school_pk):
        self.ensure_manager_access()
        year = self.get_academic_year()
        timetable = Timetable.objects.filter(school=self.get_school(), academic_year=year).first()
        if not timetable:
            raise serializers.ValidationError({"timetable": "Générez d’abord un emploi du temps."})
        if timetable.is_validated:
            raise serializers.ValidationError({"timetable": "Cet emploi du temps est déjà validé."})
        timetable.status = Timetable.Status.VALIDATED
        timetable.validated_at = timezone.now()
        timetable.save(update_fields=["status", "validated_at"])
        return Response(self.serialize(timetable))


class TimetableExportView(TimetableMixin, APIView):
    """Édition PDF de l'emploi du temps, par classe ou par enseignant.

    Un enseignant sans droit de gestion peut éditer le sien, et seulement le
    sien : c'est son propre planning, pas celui de l'établissement.
    """

    def get(self, request, school_pk):
        year = self.get_academic_year()
        timetable = Timetable.objects.filter(school=self.get_school(), academic_year=year).first()
        if not timetable:
            raise serializers.ValidationError({"timetable": "Aucun emploi du temps n’a été généré."})

        scope = request.query_params.get("scope", "classes")
        if scope == "teachers":
            return self.export_teachers(request, timetable, year)
        return self.export_classes(request, timetable, year)

    def requested_ids(self, request):
        """`?ids=3,7,12` — vide signifie « tout »."""
        raw = (request.query_params.get("ids") or "").strip()
        if not raw:
            return None
        return [int(value) for value in raw.split(",") if value.strip().isdigit()]

    def export_classes(self, request, timetable, year):
        self.ensure_manager_access()
        classes = SchoolClass.objects.filter(academic_year=year).select_related("level")
        identifiers = self.requested_ids(request)
        if identifiers is not None:
            classes = classes.filter(id__in=identifiers)
        classes = list(classes.order_by("level__order", "group"))
        if not classes:
            raise serializers.ValidationError({"classes": "Aucune classe ne correspond à la sélection."})

        content = class_timetables_pdf(timetable, classes)
        name = classes[0].group if len(classes) == 1 else f"{len(classes)}-classes"
        return self.as_attachment(content, f"emploi-du-temps-{slugify(name)}.pdf")

    def export_teachers(self, request, timetable, year):
        identifiers = self.requested_ids(request)
        teachers = CustomUser.objects.filter(
            id__in=timetable.slots.exclude(teacher=None).values_list("teacher_id", flat=True),
        )
        if identifiers is not None:
            teachers = teachers.filter(id__in=identifiers)

        # Hors gestionnaires, chacun n'a accès qu'à son propre emploi du temps.
        if not self.can_manage():
            teachers = teachers.filter(id=request.user.id)
            if not teachers.exists():
                raise serializers.ValidationError(
                    {"permission": "Vous ne pouvez éditer que votre propre emploi du temps."}
                )

        teachers = list(teachers.order_by("last_name", "first_name"))
        if not teachers:
            raise serializers.ValidationError({"teachers": "Aucun enseignant ne correspond à la sélection."})

        content = teacher_timetables_pdf(timetable, teachers)
        name = teachers[0].get_full_name() if len(teachers) == 1 else f"{len(teachers)}-enseignants"
        return self.as_attachment(content, f"emploi-du-temps-{slugify(name)}.pdf")

    def as_attachment(self, content, filename):
        response = HttpResponse(content, content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


class MyTimetableView(TimetableMixin, APIView):
    """Emploi du temps consulté depuis l'espace personnel ou le calendrier.

    Sans droit de gestion, chacun ne voit que ses propres cours. Les
    gestionnaires peuvent demander la vue générale avec `?scope=all`.
    """

    def get(self, request, school_pk):
        year = self.get_academic_year()
        timetable = Timetable.objects.filter(school=self.get_school(), academic_year=year).first()
        if not timetable:
            return Response({"timetable": None, "slots": [], "periods": [], "scope": "mine"})

        slots = timetable.slots.select_related(
            "school_class", "school_class__level", "class_subject__subject", "teacher",
        ).order_by("day", "start_time")

        # La vue générale reste réservée aux gestionnaires ; les autres
        # retombent sur leurs propres cours plutôt que sur une erreur.
        scope = "all" if request.query_params.get("scope") == "all" and self.can_manage() else "mine"
        if scope == "mine":
            slots = slots.filter(teacher=request.user)

        return Response({
            "timetable": {
                "id": timetable.id,
                "is_validated": timetable.is_validated,
                "days_per_week": timetable.days_per_week,
            },
            "scope": scope,
            "can_manage": self.can_manage(),
            "periods": [{
                "id": period.id,
                "label": period.label,
                "kind": period.kind,
                "start_time": period.start_time,
                "end_time": period.end_time,
                "order": period.order,
            } for period in timetable.periods.all()],
            "slots": [{
                "id": slot.id,
                "class_id": slot.school_class_id,
                "class_name": slot.school_class.group,
                "level": slot.school_class.level.name,
                "subject": slot.class_subject.subject.name,
                "teacher": slot.teacher.get_full_name() if slot.teacher else None,
                "teacher_id": slot.teacher_id,
                "day": slot.day,
                "day_label": DAY_LABELS.get(slot.day, ""),
                "start_time": slot.start_time,
                "end_time": slot.end_time,
            } for slot in slots],
        })


class ReportCardView(GradeMixin, APIView):
    """Bulletins d'une classe pour une session : moyennes par matière, générale et rang."""

    def get(self, request, school_pk, session_pk, class_pk):
        self.ensure_grade_access()
        year = self.get_academic_year()
        try:
            session = AcademicSession.objects.get(pk=session_pk, academic_year=year)
            school_class = SchoolClass.objects.select_related("level").get(
                pk=class_pk, academic_year=year, school=self.get_school(),
            )
        except (AcademicSession.DoesNotExist, SchoolClass.DoesNotExist):
            raise serializers.ValidationError({"context": "Session ou classe invalide."})
        if not session.classes.filter(pk=school_class.pk).exists():
            raise serializers.ValidationError({"session": "Cette classe n’appartient pas à la session."})
        try:
            scheme = GradeScheme.objects.prefetch_related("lines", "groups__lines").get(session=session)
        except GradeScheme.DoesNotExist:
            raise serializers.ValidationError(
                {"configuration": "Les lignes de notes ne sont pas configurées pour cette session."}
            )

        configurations = list(
            ClassSubject.objects.filter(school_class=school_class).select_related("subject")
        )
        enrollments = list(
            StudentEnrollment.objects.filter(
                school_class=school_class, status=StudentEnrollment.Status.ACTIVE,
            ).select_related("student").order_by("student__last_name", "student__first_name")
        )
        entries = GradeEntry.objects.filter(
            line__scheme=scheme, class_subject__school_class=school_class,
        )
        entry_map = {}
        for entry in entries:
            entry_map.setdefault((entry.enrollment_id, entry.class_subject_id), {})[entry.line_id] = entry.score

        lines = list(scheme.lines.all())
        students = []
        for enrollment in enrollments:
            subject_rows = []
            weighted_total = Decimal("0")
            coefficient_total = Decimal("0")
            for configuration in configurations:
                scores = entry_map.get((enrollment.id, configuration.id), {})
                average = (
                    GradeSheetView.calculate_average(scheme, scores)
                    if len(scores) == len(lines) and lines else None
                )
                if average is not None:
                    weighted_total += average * configuration.coefficient
                    coefficient_total += configuration.coefficient
                subject_rows.append({
                    "class_subject_id": configuration.id,
                    "subject": configuration.subject.name,
                    "coefficient": configuration.coefficient,
                    "weekly_hours": configuration.weekly_hours,
                    "average": average.quantize(Decimal("0.01")) if average is not None else None,
                })
            general = (weighted_total / coefficient_total).quantize(Decimal("0.01")) if coefficient_total else None
            students.append({
                "enrollment_id": enrollment.id,
                "matricule": enrollment.enrollment_number,
                "student_name": enrollment.student.get_full_name(),
                "subjects": subject_rows,
                "general_average": general,
                "rank": None,
            })

        # Rang : à moyenne égale, même rang ; les rangs suivants sont décalés.
        ranked = sorted(
            [student for student in students if student["general_average"] is not None],
            key=lambda student: student["general_average"], reverse=True,
        )
        previous_average = None
        previous_rank = 0
        for position, student in enumerate(ranked, start=1):
            if student["general_average"] == previous_average:
                student["rank"] = previous_rank
            else:
                student["rank"] = position
                previous_rank = position
                previous_average = student["general_average"]

        graded = [student["general_average"] for student in students if student["general_average"] is not None]
        return Response({
            "session": {"id": session.id, "name": session.name, "label": session.label},
            "school_class": {
                "id": school_class.id,
                "name": " ".join(part for part in (school_class.level.name, school_class.series, school_class.group) if part),
                "group": school_class.group,
            },
            "students": students,
            "statistics": {
                "students_total": len(students),
                "students_graded": len(graded),
                "class_average": (sum(graded) / len(graded)).quantize(Decimal("0.01")) if graded else None,
                "highest": max(graded).quantize(Decimal("0.01")) if graded else None,
                "lowest": min(graded).quantize(Decimal("0.01")) if graded else None,
                "pass_count": sum(1 for average in graded if average >= Decimal("10")),
            },
        })


# Champs de l'école imprimés en tête du bulletin. Une seule liste : la vue de
# paramétrage et l'export PDF y puisent, sans risque d'en oublier un.
REPORT_HEADER_FIELDS = (
    "country", "country_motto", "ministry", "cabinet", "general_secretariat",
    "education_direction", "direction_city", "inspection",
    "motto", "phone", "postal_box", "city",
)


def report_header(school):
    payload = {field: getattr(school, field) for field in REPORT_HEADER_FIELDS}
    payload["name"] = school.name
    # Repris par le filigrane, qui imprime au choix le nom ou le code.
    payload["code"] = school.code
    # Le paramétrage grise le filigrane « logo » tant qu'aucun logo n'est chargé.
    payload["logo"] = bool(school.logo)
    return payload


class ReportCardSettingsView(GradeMixin, APIView):
    """Mise en forme des bulletins et seuils d'appréciation, par école."""

    FIELDS = (
        "show_score_detail", "group_by_category", "show_rank", "show_teacher",
        "show_appreciation", "show_class_statistics",
        "show_principal_name", "show_censor_name", "show_founder_name",
    )

    def serialize(self, school):
        configuration = settings_for(school)
        return {
            **{field: getattr(configuration, field) for field in self.FIELDS},
            "template": configuration.template,
            "templates": [
                {"value": value, "label": label}
                for value, label in ReportCardSettings.Template.choices
            ],
            # Les filigranes se cumulent : c'est une liste, pas un choix unique.
            "watermarks": configuration.active_watermarks,
            "watermark_density": configuration.watermark_density,
            "watermark_source": configuration.watermark_source,
            # Tant que c'est faux, les valeurs affichées sont celles d'usage.
            "is_configured": configuration.configured_at is not None,
            "watermark_choices": [
                {"value": value, "label": label}
                for value, label in ReportCardSettings.Watermark.choices
                if value != ReportCardSettings.Watermark.NONE
            ],
            "watermark_densities": [
                {"value": value, "label": label}
                for value, label in ReportCardSettings.WatermarkDensity.choices
            ],
            "watermark_sources": [
                {"value": value, "label": label}
                for value, label in ReportCardSettings.WatermarkSource.choices
            ],
            "council_note": configuration.council_note,
            # Qui signera, d'après les rôles de l'école : le paramétrage
            # montre le nom qu'il imprimera, ou le poste resté vacant.
            "signatories": role_holders(school),
            "appreciations": [
                {"id": row.id, "label": row.label, "minimum": str(row.minimum)}
                for row in appreciations_for(school)
            ],
            "category_orders": [
                {
                    "id": rule.id,
                    "name": rule.name,
                    "scope": rule.scope,
                    "scope_label": rule.get_scope_display(),
                    "stage": rule.stage,
                    "series": rule.series,
                    "classes": [item.id for item in rule.classes.all()],
                    "class_names": [item.group for item in rule.classes.all()],
                    "categories": rule.categories or [],
                }
                for rule in SubjectCategoryOrder.objects.filter(school=school).prefetch_related("classes")
            ],
            "categories": [
                {"id": row.id, "name": row.name}
                for row in SubjectCategory.objects.filter(school=school)
            ],
            "stages": [
                {"value": value, "label": label}
                for value, label in SchoolLevel.Stage.choices
            ],
            # `distinct()` ne suffit pas : l'ordonnancement par défaut de
            # SchoolClass ajoute ses colonnes au SELECT et rend les lignes
            # uniques. On dédoublonne donc en Python.
            "series": sorted(set(
                SchoolClass.objects.filter(school=school, is_active=True)
                .exclude(series="").values_list("series", flat=True)
            )),
            "school": report_header(school),
            "can_configure": self.can_configure_grades(),
        }

    def get(self, request, school_pk):
        self.ensure_grade_access()
        return Response(self.serialize(self.get_school()))

    @transaction.atomic
    def put(self, request, school_pk):
        if not self.can_configure_grades():
            raise serializers.ValidationError(
                {"permission": "Seule la direction de l’établissement peut configurer les bulletins."}
            )
        school = self.get_school()
        configuration = settings_for(school)
        for field in self.FIELDS:
            if field in request.data:
                setattr(configuration, field, bool(request.data[field]))
        if "council_note" in request.data:
            configuration.council_note = str(request.data["council_note"])[:2000]
        # Champs à choix fermé : une valeur inconnue est refusée plutôt
        # qu'enregistrée telle quelle, sinon l'édition PDF retomberait
        # silencieusement sur son défaut.
        choice_fields = (
            ("template", ReportCardSettings.Template, "Modèle de bulletin inconnu."),
            ("watermark_density", ReportCardSettings.WatermarkDensity, "Densité de mosaïque inconnue."),
            ("watermark_source", ReportCardSettings.WatermarkSource, "Texte de filigrane inconnu."),
        )
        for field, choices, message in choice_fields:
            if field in request.data:
                value = str(request.data[field])
                if value not in dict(choices.choices):
                    raise serializers.ValidationError({field: message})
                setattr(configuration, field, value)

        # Les filigranes se combinent : on enregistre une liste, dédoublonnée
        # et purgée de « aucun », qui ne veut rien dire à côté d'un autre.
        if "watermarks" in request.data:
            chosen = request.data["watermarks"]
            if not isinstance(chosen, list):
                raise serializers.ValidationError(
                    {"watermarks": "Les filigranes doivent être envoyés sous forme de liste."}
                )
            known = dict(ReportCardSettings.Watermark.choices)
            kinds = []
            for item in chosen:
                value = str(item)
                if value not in known:
                    raise serializers.ValidationError({"watermarks": "Filigrane inconnu."})
                if value != ReportCardSettings.Watermark.NONE and value not in kinds:
                    kinds.append(value)
            configuration.watermarks = kinds
            # L'ancien champ suit, pour les lectures qui n'ont pas migré.
            configuration.watermark = kinds[0] if kinds else ReportCardSettings.Watermark.NONE
        # Marque le paramétrage comme réglé à la main : à partir d'ici, aucune
        # valeur d'usage ne vient plus le compléter dans le dos de l'école.
        if configuration.configured_at is None:
            configuration.configured_at = timezone.now()
        configuration.save()

        # Coordonnées imprimées en tête du bulletin.
        details = request.data.get("school") or {}
        touched = [field for field in REPORT_HEADER_FIELDS if field in details]
        for field in touched:
            # Chaque champ a sa propre longueur : tronquer à une valeur commune
            # écourterait les uns et ferait échouer l'écriture des autres.
            limit = School._meta.get_field(field).max_length
            setattr(school, field, str(details[field]).strip()[:limit])
        if touched:
            school.save(update_fields=touched)

        appreciations = request.data.get("appreciations")
        if appreciations is not None:
            rows = []
            for item in appreciations:
                label = str(item.get("label", "")).strip()
                if not label:
                    continue
                try:
                    minimum = Decimal(str(item.get("minimum")))
                except (TypeError, ArithmeticError):
                    raise serializers.ValidationError(
                        {"appreciations": f"Seuil invalide pour « {label} »."}
                    )
                if not Decimal("0") <= minimum <= Decimal("20"):
                    raise serializers.ValidationError(
                        {"appreciations": f"Le seuil de « {label} » doit être compris entre 0 et 20."}
                    )
                rows.append((label, minimum))

            labels = [label for label, _ in rows]
            if len(labels) != len(set(labels)):
                raise serializers.ValidationError(
                    {"appreciations": "Deux appréciations portent le même libellé."}
                )
            ReportCardAppreciation.objects.filter(school=school).delete()
            ReportCardAppreciation.objects.bulk_create([
                ReportCardAppreciation(school=school, label=label, minimum=minimum)
                for label, minimum in rows
            ])

        category_orders = request.data.get("category_orders")
        if category_orders is not None:
            valid_categories = set(
                SubjectCategory.objects.filter(school=school).values_list("name", flat=True)
            )
            valid_classes = set(
                SchoolClass.objects.filter(school=school).values_list("id", flat=True)
            )
            valid_stages = {value for value, _ in SchoolLevel.Stage.choices}

            prepared = []
            for item in category_orders:
                scope = item.get("scope", SubjectCategoryOrder.Scope.SCHOOL)
                if scope not in dict(SubjectCategoryOrder.Scope.choices):
                    raise serializers.ValidationError({"category_orders": "Portée inconnue."})

                stage = str(item.get("stage") or "")
                if scope == SubjectCategoryOrder.Scope.STAGE and stage not in valid_stages:
                    raise serializers.ValidationError({"category_orders": "Cycle invalide."})

                series = str(item.get("series") or "").strip()
                if scope == SubjectCategoryOrder.Scope.SERIES and not series:
                    raise serializers.ValidationError({"category_orders": "Indiquez une série."})

                members = [
                    class_id for class_id in dict.fromkeys(item.get("classes") or [])
                    if class_id in valid_classes
                ]
                if scope == SubjectCategoryOrder.Scope.CLASSES and not members:
                    raise serializers.ValidationError(
                        {"category_orders": "Choisissez au moins une classe."}
                    )

                # Un type disparu de l'école ne doit pas figer un ordre obsolète.
                ordered = [
                    str(name) for name in (item.get("categories") or [])
                    if str(name) in valid_categories
                ]
                prepared.append((item.get("name", ""), scope, stage, series, members, ordered))

            SubjectCategoryOrder.objects.filter(school=school).delete()
            for name, scope, stage, series, members, ordered in prepared:
                rule = SubjectCategoryOrder.objects.create(
                    school=school, name=str(name)[:80], scope=scope,
                    stage=stage if scope == SubjectCategoryOrder.Scope.STAGE else "",
                    series=series if scope == SubjectCategoryOrder.Scope.SERIES else "",
                    categories=ordered,
                )
                if scope == SubjectCategoryOrder.Scope.CLASSES:
                    rule.classes.set(members)

        return Response(self.serialize(school))


class ReportCardGenerationView(GradeMixin, APIView):
    """Génère les bulletins d'une session : établissement, classe ou élève.

    Tant que la session reste ouverte, tout peut être régénéré : un barème
    corrigé ou une note rectifiée doit pouvoir se répercuter sur les bulletins
    déjà édités. La clôture, elle, fige l'ensemble — plus aucune génération
    n'est alors acceptée, quelle que soit la portée.
    """

    def get_session(self, session_pk):
        year = self.get_academic_year()
        try:
            return AcademicSession.objects.get(pk=session_pk, academic_year=year)
        except AcademicSession.DoesNotExist:
            raise serializers.ValidationError({"session": "Session invalide."})

    def get(self, request, school_pk, session_pk):
        """État de génération : combien de bulletins, quelles classes."""
        self.ensure_grade_access()
        session = self.get_session(session_pk)
        school = self.get_school()
        cards = ReportCard.objects.filter(session=session, school_class__school=school)
        per_class = {}
        for card in cards.select_related("school_class", "school_class__level"):
            entry = per_class.setdefault(card.school_class_id, {
                "id": card.school_class_id,
                "name": card.school_class.group,
                "level": card.school_class.level.name,
                "count": 0,
                "generated_at": None,
            })
            entry["count"] += 1
            stamp = card.generated_at.isoformat()
            if entry["generated_at"] is None or stamp > entry["generated_at"]:
                entry["generated_at"] = stamp

        return Response({
            "session": {"id": session.id, "name": session.name, "label": session.label},
            "generated": cards.exists(),
            "total": cards.count(),
            "classes": sorted(per_class.values(), key=lambda row: row["name"]),
            "can_configure": self.can_configure_grades(),
        })

    def post(self, request, school_pk, session_pk):
        if not self.can_configure_grades():
            raise serializers.ValidationError(
                {"permission": "Seuls le propriétaire, l’administrateur, le censeur ou le proviseur peuvent générer les bulletins."}
            )
        session = self.get_session(session_pk)
        school = self.get_school()
        if session.is_closed:
            raise serializers.ValidationError(
                {"session": "Cette session est clôturée : les bulletins ne peuvent plus être régénérés."}
            )
        scope = request.data.get("scope", "school")
        issued_on = self.get_issued_on(request.data.get("issued_on"))

        try:
            if scope == "school":
                already = ReportCard.objects.filter(session=session, school_class__school=school).exists()
                written, classes = generate_school(session, school, user=request.user, issued_on=issued_on)
                verb = "régénérés" if already else "générés"
                message = f"{written} bulletin(s) {verb} pour {classes} classe(s)."

            elif scope == "class":
                school_class = self.get_class(session, request.data.get("school_class"))
                written = generate_class(session, school_class, user=request.user, issued_on=issued_on)
                message = f"{written} bulletin(s) régénérés pour {school_class.group}."

            elif scope == "student":
                enrollment = self.get_enrollment(session, request.data.get("enrollment"))
                written = generate_class(
                    session, enrollment.school_class, user=request.user,
                    enrollment_ids={enrollment.id}, issued_on=issued_on,
                )
                message = f"Bulletin régénéré pour {enrollment.student.get_full_name()}."

            else:
                raise serializers.ValidationError({"scope": "Portée inconnue."})
        except ValueError as issue:
            raise serializers.ValidationError({"configuration": str(issue)})

        return Response({"detail": message, "written": written})

    def get_issued_on(self, value):
        """Date d'établissement portée par les bulletins, « fait le … ».

        Absente, on prend le jour de la génération : le bulletin imprime une
        date plutôt qu'une ligne de pointillés.
        """
        if value in (None, ""):
            return timezone.localdate()
        parsed = parse_date(str(value))
        if parsed is None:
            raise serializers.ValidationError(
                {"issued_on": "Date d’établissement invalide (attendu : AAAA-MM-JJ)."}
            )
        return parsed

    def get_class(self, session, class_id):
        try:
            school_class = SchoolClass.objects.select_related("level").get(
                pk=class_id, school=self.get_school(), academic_year=session.academic_year,
            )
        except (SchoolClass.DoesNotExist, ValueError, TypeError):
            raise serializers.ValidationError({"school_class": "Classe invalide."})
        if not session.classes.filter(pk=school_class.pk).exists():
            raise serializers.ValidationError(
                {"school_class": "Cette classe n’appartient pas à la session."}
            )
        return school_class

    def get_enrollment(self, session, enrollment_id):
        try:
            enrollment = StudentEnrollment.objects.select_related(
                "student", "school_class",
            ).get(pk=enrollment_id, school_class__school=self.get_school())
        except (StudentEnrollment.DoesNotExist, ValueError, TypeError):
            raise serializers.ValidationError({"enrollment": "Élève invalide."})
        if not session.classes.filter(pk=enrollment.school_class_id).exists():
            raise serializers.ValidationError(
                {"enrollment": "Cet élève n’appartient pas à une classe de la session."}
            )
        return enrollment


class MySignatureView(APIView):
    """Signature manuscrite de l'utilisateur connecté.

    Elle est tracée à l'écran puis envoyée en PNG encodé : c'est la même
    signature que celle du bas des bulletins, chacun enregistre la sienne et
    personne ne signe pour un autre.
    """

    # Une signature tracée à la main tient largement dedans ; au-delà, c'est
    # une image importée d'ailleurs, que ce point d'entrée n'a pas à recevoir.
    MAX_BYTES = 512 * 1024
    PREFIX = "data:image/png;base64,"

    def serialize(self, user):
        return {
            "signature": user.signature.url if user.signature else None,
            "has_signature": bool(user.signature),
        }

    def get(self, request):
        return Response(self.serialize(request.user))

    def put(self, request):
        image = request.data.get("image")
        if not isinstance(image, str) or not image.startswith(self.PREFIX):
            raise serializers.ValidationError(
                {"image": "Envoyez la signature en PNG encodé (data:image/png;base64,…)."}
            )
        try:
            content = base64.b64decode(image[len(self.PREFIX):], validate=True)
        except (binascii.Error, ValueError):
            raise serializers.ValidationError({"image": "Image illisible."})
        if not content:
            raise serializers.ValidationError({"image": "Signature vide."})
        if len(content) > self.MAX_BYTES:
            raise serializers.ValidationError(
                {"image": "Signature trop lourde : 512 Ko au maximum."}
            )
        # Un PNG commence toujours par cette signature de fichier ; sans ce
        # contrôle, n'importe quel binaire passerait pour une image.
        if not content.startswith(b"\x89PNG\r\n\x1a\n"):
            raise serializers.ValidationError({"image": "Format inattendu : PNG attendu."})

        content = self.trim(content)
        user = request.user
        # L'ancienne signature part avec la nouvelle : sans cela, chaque
        # enregistrement laisserait un fichier de plus sur le disque.
        user.signature.delete(save=False)
        user.signature.save(f"signature-{user.pk}.png", ContentFile(content), save=True)
        return Response(self.serialize(user))

    @staticmethod
    def trim(content):
        """Rogne le vide autour du tracé et rend le fond transparent.

        On signe rarement au milieu du cadre : sans ce recadrage, le bulletin
        réduirait toute la zone de dessin — marges comprises — et la signature
        y paraîtrait minuscule et décentrée. Un fond blanc masquerait en outre
        le filigrane, d'où la transparence.

        L'image d'origine est renvoyée telle quelle si elle résiste : mieux
        vaut une signature mal cadrée que pas de signature du tout.
        """
        try:
            image = PillowImage.open(io.BytesIO(content)).convert("RGBA")
            # Signature envoyée sur fond opaque : on efface le blanc pour ne
            # garder que l'encre, sinon le cadre entier passerait pour du tracé.
            if image.getchannel("A").getextrema()[0] == 255:
                ink = image.convert("L").point(lambda level: 0 if level > 240 else 255)
                image.putalpha(ink)
            box = image.getbbox()
            if box is None:
                return content
            # Une marge fine évite que le trait ne touche le bord du cadre.
            margin = max(4, min(image.width, image.height) // 60)
            cropped = image.crop((
                max(0, box[0] - margin), max(0, box[1] - margin),
                min(image.width, box[2] + margin), min(image.height, box[3] + margin),
            ))
            buffer = io.BytesIO()
            cropped.save(buffer, format="PNG", optimize=True)
            return buffer.getvalue()
        except Exception:
            return content

    def delete(self, request):
        request.user.signature.delete(save=True)
        return Response(self.serialize(request.user))


class PromotionThresholdView(GradeMixin, APIView):
    """Moyennes exigées pour passer, niveau par niveau.

    Le seuil se règle par niveau et non par école : un établissement peut
    exiger 10 au lycée et se montrer plus souple en primaire. Sur un niveau
    d'examen, le seuil porte sur la note de l'examen officiel.
    """

    def levels(self):
        return SchoolLevel.objects.filter(school=self.get_school()).order_by("order")

    def serialize(self, level):
        return {
            "id": level.id,
            "name": level.name,
            "stage": level.stage,
            "order": level.order,
            "passing_average": str(level.passing_average),
            "is_exam_level": level.is_exam_level,
            "exam_name": level.exam_name,
        }

    def get(self, request, school_pk):
        self.ensure_grade_access()
        return Response({
            "can_configure": self.can_configure_grades(),
            "levels": [self.serialize(level) for level in self.levels()],
        })

    def put(self, request, school_pk):
        self.ensure_grade_configurator()
        rows = request.data.get("levels")
        if not isinstance(rows, list):
            raise serializers.ValidationError({"levels": "Envoyez la liste des niveaux."})

        by_id = {level.id: level for level in self.levels()}
        updated = []
        for row in rows:
            level = by_id.get(row.get("id"))
            if level is None:
                raise serializers.ValidationError({"levels": "Niveau inconnu dans cette école."})
            level.passing_average = self.get_average(row.get("passing_average"))
            level.is_exam_level = bool(row.get("is_exam_level"))
            # Un niveau qui n'est plus d'examen ne garde pas le nom de l'ancien.
            level.exam_name = (row.get("exam_name") or "").strip() if level.is_exam_level else ""
            updated.append(level)

        SchoolLevel.objects.bulk_update(
            updated, ["passing_average", "is_exam_level", "exam_name"],
        )
        return Response({"levels": [self.serialize(level) for level in self.levels()]})

    @staticmethod
    def get_average(value):
        """Une moyenne sur 20, refusée hors bornes plutôt que silencieusement rognée."""
        try:
            average = Decimal(str(value))
        except (InvalidOperation, TypeError, ValueError):
            raise serializers.ValidationError({"passing_average": "Moyenne invalide."})
        if not Decimal("0") <= average <= Decimal("20"):
            raise serializers.ValidationError(
                {"passing_average": "La moyenne de passage se situe entre 0 et 20."}
            )
        return average.quantize(Decimal("0.01"))


class YearEndDecisionView(GradeMixin, APIView):
    """Décisions de fin d'année d'une classe, et saisie des notes d'examen.

    N'a de sens que sur la dernière session de l'année : c'est elle qui porte
    la moyenne annuelle, donc la décision de passage.
    """

    def get_context(self, session_pk, class_id):
        year = self.get_academic_year()
        try:
            session = AcademicSession.objects.get(pk=session_pk, academic_year=year)
            school_class = SchoolClass.objects.select_related("level").get(
                pk=class_id, academic_year=year, school=self.get_school(),
            )
        except (AcademicSession.DoesNotExist, SchoolClass.DoesNotExist, TypeError, ValueError):
            raise serializers.ValidationError({"context": "Session ou classe invalide."})
        if not session.classes.filter(pk=school_class.pk).exists():
            raise serializers.ValidationError(
                {"session": "Cette classe n’appartient pas à la session."}
            )
        if not session.is_final:
            raise serializers.ValidationError({"session": (
                "Les décisions de fin d'année ne se prennent que sur la dernière "
                "session de l'année."
            )})
        return session, school_class

    def rows(self, session, school_class):
        """Un élève par ligne : moyenne annuelle, note d'examen, décision."""
        cards = list(
            ReportCard.objects.filter(session=session, school_class=school_class)
            .select_related("enrollment__student")
            .order_by("enrollment__student__last_name", "enrollment__student__first_name")
        )
        history = term_history(
            session, school_class, [card.enrollment_id for card in cards],
        )
        level = school_class.level
        rows = []
        for card in cards:
            student = student_history(history, card.enrollment_id)
            decision = promotion_decision(
                level, student["annual_average"], card.exam_average,
                gender=card.enrollment.student.gender,
            )
            rows.append({
                "enrollment": card.enrollment_id,
                "matricule": card.enrollment.enrollment_number,
                "student_name": card.enrollment.student.get_full_name(),
                "annual_average": student["annual_average"],
                "exam_average": str(card.exam_average) if card.exam_average is not None else None,
                "decision": decision,
            })
        return rows

    def get(self, request, school_pk, session_pk):
        self.ensure_grade_access()
        session, school_class = self.get_context(session_pk, request.query_params.get("school_class"))
        level = school_class.level
        return Response({
            "can_configure": self.can_configure_grades(),
            "is_closed": session.is_closed,
            "level": {
                "id": level.id,
                "name": level.name,
                "is_exam_level": level.is_exam_level,
                "exam_name": level.exam_name,
                "passing_average": str(level.passing_average),
            },
            "students": self.rows(session, school_class),
        })

    def post(self, request, school_pk, session_pk):
        """Enregistre les notes d'examen saisies pour la classe."""
        self.ensure_grade_configurator()
        session, school_class = self.get_context(session_pk, request.data.get("school_class"))
        if session.is_closed:
            raise serializers.ValidationError(
                {"session": "Cette session est clôturée : les décisions sont figées."}
            )
        if not school_class.level.is_exam_level:
            raise serializers.ValidationError({"school_class": (
                "Ce niveau n'est pas un niveau d'examen : sa décision suit la "
                "moyenne annuelle, il n'y a pas de note d'examen à saisir."
            )})

        results = request.data.get("results")
        if not isinstance(results, list):
            raise serializers.ValidationError({"results": "Envoyez la liste des résultats."})

        cards = {
            card.enrollment_id: card for card in
            ReportCard.objects.filter(session=session, school_class=school_class)
        }
        touched = []
        for result in results:
            card = cards.get(result.get("enrollment"))
            if card is None:
                raise serializers.ValidationError(
                    {"results": "Élève sans bulletin dans cette classe."}
                )
            card.exam_average = self.get_exam_average(result.get("exam_average"))
            touched.append(card)
        ReportCard.objects.bulk_update(touched, ["exam_average"])

        return Response({
            "detail": f"{len(touched)} résultat(s) d'examen enregistré(s).",
            "students": self.rows(session, school_class),
        })

    @staticmethod
    def get_exam_average(value):
        """Note d'examen sur 20 ; vide efface le résultat déjà saisi."""
        if value in (None, ""):
            return None
        try:
            average = Decimal(str(value))
        except (InvalidOperation, TypeError, ValueError):
            raise serializers.ValidationError({"exam_average": "Note d'examen invalide."})
        if not Decimal("0") <= average <= Decimal("20"):
            raise serializers.ValidationError(
                {"exam_average": "La note d'examen se situe entre 0 et 20."}
            )
        return average.quantize(Decimal("0.01"))


class ReportCardExportView(GradeMixin, APIView):
    """Édite en PDF les bulletins figés : établissement, classe ou élève."""

    def get(self, request, school_pk, session_pk):
        self.ensure_grade_access()
        year = self.get_academic_year()
        school = self.get_school()
        try:
            session = AcademicSession.objects.get(pk=session_pk, academic_year=year)
        except AcademicSession.DoesNotExist:
            raise serializers.ValidationError({"session": "Session invalide."})

        cards = ReportCard.objects.filter(
            session=session, school_class__school=school,
        ).select_related(
            "school_class", "school_class__level", "school_class__homeroom_teacher",
            "enrollment__student",
        )

        scope = request.query_params.get("scope", "school")
        if scope == "class":
            cards = cards.filter(school_class_id=request.query_params.get("school_class"))
        elif scope == "student":
            cards = cards.filter(enrollment_id=request.query_params.get("enrollment"))

        cards = list(cards.order_by("school_class__group", "rank"))
        if not cards:
            raise serializers.ValidationError(
                {"report_cards": "Aucun bulletin généré pour cette sélection."}
            )

        configuration = settings_for(school)
        options = {
            "template": configuration.template,
            "watermarks": configuration.active_watermarks,
            "watermark_density": configuration.watermark_density,
            "watermark_source": configuration.watermark_source,
            "show_score_detail": configuration.show_score_detail,
            "show_rank": configuration.show_rank,
            "show_teacher": configuration.show_teacher,
            "show_appreciation": configuration.show_appreciation,
            "show_class_statistics": configuration.show_class_statistics,
            "council_note": configuration.council_note,
            # Signataires du pied de page, résolus depuis les rôles de l'école.
            "signatories": signatories_for(school, configuration),
            # Signature de l'enseignant de chaque matière, pour la colonne
            # « Signature » de la grille. Les identifiants de configuration
            # sont propres à une classe : les classes se cumulent sans risque.
            "teacher_signatures": {
                key: path
                for school_class in {card.school_class for card in cards}
                for key, path in teacher_signatures_for(school_class).items()
            },
        }

        # Le récapitulatif reprend les sessions déjà éditées de l'année. La
        # séquence est propre à chaque classe — une année peut mêler classes en
        # trimestres et classes en semestres — d'où un historique par classe.
        histories = {}
        for school_class in {card.school_class for card in cards}:
            enrollment_ids = [
                card.enrollment_id for card in cards
                if card.school_class_id == school_class.id
            ]
            histories[school_class.id] = term_history(session, school_class, enrollment_ids)

        payloads = []
        for card in cards:
            payload = dict(card.payload or {})
            payload["class_name"] = " ".join(
                part for part in (card.school_class.level.name, card.school_class.series,
                                  card.school_class.group) if part
            )
            payload["history"] = student_history(
                histories[card.school_class_id], card.enrollment_id,
            )
            # Date d'établissement choisie à la génération : elle remplace la
            # ligne de pointillés au pied du bulletin.
            payload["issued_on"] = date_format(card.issued_on, "DATE_FORMAT") if card.issued_on else ""
            # Décision de fin d'année, sur la dernière session seulement : le
            # passage se joue sur la moyenne annuelle, ou sur l'examen.
            payload["homeroom_signature"] = signature_path(card.school_class.homeroom_teacher)
            payload["decision"] = promotion_decision(
                card.school_class.level, payload["history"]["annual_average"], card.exam_average,
                gender=card.enrollment.student.gender,
            ) if payload["history"]["is_final"] else None
            payloads.append(payload)

        content = report_cards_pdf(
            payloads,
            school=report_header(school),
            session={"id": session.id, "name": session.name, "label": session.label},
            year_name=year.name,
            options=options,
            logo_path=school.logo.path if school.logo else None,
        )

        return self.as_attachment(content, self.download_name(scope, cards, session, year, school))

    @staticmethod
    def download_name(scope, cards, session, year, school):
        """Nom du fichier téléchargé, lisible tel quel dans un dossier.

        « Bulletin du Troisième Trimestre 2025-2026 de KODJO ABOTSI » : la
        session, l'année et ce que porte le fichier. Le pluriel suit le
        contenu — un élève donne un bulletin, une classe en donne plusieurs.
        """
        if scope == "student":
            target = cards[0].enrollment.student.get_full_name()
        elif scope == "class":
            school_class = cards[0].school_class
            target = " ".join(part for part in (
                school_class.level.name, school_class.series, school_class.group,
            ) if part)
        else:
            target = school.name
        noun = "Bulletin" if scope == "student" else "Bulletins"
        return f"{noun} du {session.name} {year.name} de {target}.pdf"

    @staticmethod
    def as_attachment(content, filename):
        """Réponse PDF téléchargée sous un nom lisible, accents compris.

        Deux formes dans l'en-tête : la version ASCII pour les clients anciens,
        et `filename*` en UTF-8, que les navigateurs préfèrent quand elle est
        là. Les caractères interdits dans un nom de fichier sont écartés, sans
        quoi le téléchargement échouerait sur certains systèmes.
        """
        clean = re.sub(r'[\\/:*?"<>|\r\n\t]', " ", filename)
        clean = re.sub(r"\s+", " ", clean).strip()
        ascii_name = unicodedata.normalize("NFKD", clean).encode("ascii", "ignore").decode() or "bulletins.pdf"
        response = HttpResponse(content, content_type="application/pdf")
        response["Content-Disposition"] = (
            f'attachment; filename="{ascii_name}"; '
            f"filename*=UTF-8''{quote(clean)}"
        )
        return response


def save_attachments(message, files):
    """Attache les fichiers reçus à un message, après contrôle.

    Refuse plutôt qu'ignorer : un fichier trop lourd ou d'un type non prévu
    doit se voir signalé à l'expéditeur, pas disparaître en silence.
    """
    saved = []
    for upload in files:
        if upload.size > MessageAttachment.MAX_SIZE:
            limit = MessageAttachment.MAX_SIZE // (1024 * 1024)
            raise serializers.ValidationError(
                {"attachments": f"« {upload.name} » dépasse {limit} Mo."}
            )
        kind = MessageAttachment.resolve_kind(
            getattr(upload, "content_type", ""), upload.name,
        )
        if kind is None:
            raise serializers.ValidationError(
                {"attachments": f"Type de fichier non autorisé : « {upload.name} »."}
            )
        saved.append(MessageAttachment.objects.create(
            message=message,
            file=upload,
            kind=kind,
            original_name=upload.name[:255],
            content_type=(getattr(upload, "content_type", "") or "")[:120],
            size=upload.size,
        ))
    return saved


def attachment_duration(request):
    """Durée déclarée par le client pour un vocal ou une vidéo, si fournie."""
    raw = request.data.get("duration_seconds")
    try:
        value = int(raw)
    except (TypeError, ValueError):
        return None
    return value if 0 < value < 60 * 60 * 6 else None


class CommunicationMixin(SchoolScopedMixin):
    """Périmètre commun aux annonces et à la messagerie."""

    def get_academic_year(self):
        year_id = self.request.headers.get("X-Academic-Year-ID")
        if not year_id:
            raise serializers.ValidationError({"academic_year": "Sélectionnez une année académique."})
        try:
            return AcademicYear.objects.get(pk=year_id, school=self.get_school())
        except (AcademicYear.DoesNotExist, ValueError):
            raise serializers.ValidationError({"academic_year": "Année académique invalide."})

    def membership_role(self):
        """Rôle tenu dans cette école, à défaut le rôle du compte."""
        school = self.get_school()
        if school.owner_id == self.request.user.id:
            return CustomUser.Role.OWNER
        membership = school.memberships.filter(
            user=self.request.user, is_active=True,
        ).first()
        return membership.role if membership else self.request.user.role

    def can_publish(self):
        return (
            self.request.user.is_superuser
            or self.membership_role() in DIRECTION_ROLES
        )

    def ensure_publisher(self):
        if not self.can_publish():
            raise serializers.ValidationError(
                {"permission": "Seuls le propriétaire, l’administrateur, le censeur "
                               "ou le proviseur peuvent publier une annonce."}
            )

    def taught_class_ids(self, year):
        return set(
            TeacherClassAssignment.objects.filter(
                teacher=self.request.user, school=self.get_school(), academic_year=year,
            ).values_list("school_class_id", flat=True)
        )

    def guarded_class_ids(self, year):
        """Classes des enfants dont l'utilisateur est tuteur.

        Un parent n'a pas d'appartenance à l'école : c'est ce lien de tutelle
        qui le rattache aux classes de ses enfants.
        """
        return set(
            StudentEnrollment.objects.filter(
                guardian=self.request.user, school=self.get_school(),
                academic_year=year, status=StudentEnrollment.Status.ACTIVE,
            ).values_list("school_class_id", flat=True)
        )

    def visible_announcements(self, year):
        """Annonces publiées qui s'adressent à cet utilisateur."""
        role = self.membership_role()
        today = timezone.localdate()

        queryset = Announcement.objects.filter(
            school=self.get_school(), academic_year=year, is_published=True,
        ).filter(
            # Une annonce expirée sort des listes ; sans date d'expiration
            # elle reste visible indéfiniment.
            Q(expires_on__isnull=True) | Q(expires_on__gte=today)
        )

        if self.can_publish():
            # La direction voit tout, y compris ce qu'elle n'a pas écrit :
            # elle doit pouvoir relire ce qui circule dans l'établissement.
            return queryset.distinct()

        is_staff_member = role not in (CustomUser.Role.PARENT, CustomUser.Role.STUDENT)
        reachable = self.taught_class_ids(year) | self.guarded_class_ids(year)

        visibility = Q(audience=Announcement.Audience.EVERYONE)
        if is_staff_member:
            visibility |= Q(audience=Announcement.Audience.STAFF)
        visibility |= Q(audience=Announcement.Audience.ROLES, roles__contains=role)
        if reachable:
            visibility |= Q(
                audience=Announcement.Audience.CLASSES, classes__id__in=reachable,
            )
        return queryset.filter(visibility).distinct()


class AnnouncementListView(CommunicationMixin, APIView):
    """Annonces visibles par l'utilisateur, et publication pour la direction."""

    def get(self, request, school_pk):
        year = self.get_academic_year()
        announcements = list(
            self.visible_announcements(year)
            .select_related("author")
            .prefetch_related("classes")[:100]
        )
        read_ids = set(
            AnnouncementRead.objects.filter(
                user=request.user, announcement__in=announcements,
            ).values_list("announcement_id", flat=True)
        )
        return Response({
            "announcements": AnnouncementSerializer(
                announcements, many=True, context={"read_ids": read_ids},
            ).data,
            "unread_count": sum(
                1 for item in announcements if item.id not in read_ids
            ),
            "can_publish": self.can_publish(),
        })

    @transaction.atomic
    def post(self, request, school_pk):
        self.ensure_publisher()
        year = self.get_academic_year()
        serializer = AnnouncementSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        classes = serializer.validated_data.pop("classes", [])
        announcement = serializer.save(
            school=self.get_school(), academic_year=year, author=request.user,
        )
        if classes:
            invalid = [
                item.group for item in classes
                if item.school_id != self.get_school().id
                or item.academic_year_id != year.id
            ]
            if invalid:
                raise serializers.ValidationError(
                    {"classes": f"Classe hors de cette école ou de cette année : {', '.join(invalid)}."}
                )
            announcement.classes.set(classes)
        return Response(
            AnnouncementSerializer(announcement, context={"read_ids": set()}).data,
            status=201,
        )


class AnnouncementDetailView(CommunicationMixin, APIView):
    """Marquer lue, ou retirer une annonce."""

    def get_announcement(self, year, pk):
        try:
            return self.visible_announcements(year).get(pk=pk)
        except (Announcement.DoesNotExist, ValueError):
            raise serializers.ValidationError({"announcement": "Annonce introuvable."})

    def post(self, request, school_pk, pk):
        """Accusé de lecture. Rejouer l'appel ne crée pas de doublon."""
        year = self.get_academic_year()
        announcement = self.get_announcement(year, pk)
        AnnouncementRead.objects.get_or_create(
            announcement=announcement, user=request.user,
        )
        return Response({"detail": "Annonce marquée comme lue."})

    def delete(self, request, school_pk, pk):
        self.ensure_publisher()
        year = self.get_academic_year()
        self.get_announcement(year, pk).delete()
        return Response(status=204)


class ConversationListView(CommunicationMixin, APIView):
    """Fils de discussion de l'utilisateur, et ouverture d'un nouveau fil."""

    def get(self, request, school_pk):
        conversations = (
            Conversation.objects
            .filter(school=self.get_school(), participants=request.user)
            .prefetch_related("participants", "messages")
            .distinct()
        )
        data = ConversationSerializer(
            conversations, many=True, context={"user_id": request.user.id},
        ).data
        return Response({
            "conversations": data,
            "unread_count": sum(item["unread_count"] for item in data),
        })

    @transaction.atomic
    def post(self, request, school_pk):
        """Ouvre un fil et y dépose le premier message."""
        school = self.get_school()
        body = str(request.data.get("body", "")).strip()
        files = request.FILES.getlist("attachments")
        if not body and not files:
            raise serializers.ValidationError({"body": "Le message ne peut pas être vide."})

        # En JSON les destinataires arrivent en liste ; en multipart — le
        # format employé dès qu'il y a une pièce jointe — ils arrivent en
        # champs répétés, à relire avec `getlist`, et sous forme de chaînes.
        if hasattr(request.data, "getlist"):
            raw_ids = request.data.getlist("participants")
        else:
            raw_ids = request.data.get("participants")
        if not isinstance(raw_ids, (list, tuple)):
            raw_ids = [raw_ids] if raw_ids not in (None, "") else []

        wanted = set()
        for value in raw_ids:
            try:
                wanted.add(int(value))
            except (TypeError, ValueError):
                raise serializers.ValidationError(
                    {"participants": "Destinataire invalide."}
                )
        if not wanted:
            raise serializers.ValidationError(
                {"participants": "Choisissez au moins un destinataire."}
            )

        people = [person for person in self.reachable_users() if person.id in wanted]
        if len(people) != len(wanted):
            raise serializers.ValidationError(
                {"participants": "Un destinataire n’est pas joignable dans cette école."}
            )

        conversation = Conversation.objects.create(
            school=school,
            subject=str(request.data.get("subject", "")).strip()[:180],
            started_by=request.user,
        )
        conversation.participants.set([request.user, *people])
        message = Message.objects.create(
            conversation=conversation, sender=request.user, body=body,
        )
        save_attachments(message, files)
        # L'expéditeur a évidemment lu son propre message.
        message.read_by.add(request.user)
        conversation.last_message_at = message.sent_at
        conversation.save(update_fields=["last_message_at"])

        return Response(
            ConversationSerializer(
                conversation, context={"user_id": request.user.id},
            ).data,
            status=201,
        )

    def reachable_users(self):
        """Personnes joignables : le personnel de l'école, et les parents.

        Le personnel se joint par son appartenance à l'école ; les parents,
        par le lien de tutelle sur une inscription — ils n'ont pas
        d'appartenance.
        """
        school = self.get_school()
        staff = Q(school_memberships__school=school, school_memberships__is_active=True)
        guardians = Q(guarded_student_enrollments__school=school)
        owner = Q(pk=school.owner_id)
        return (
            CustomUser.objects
            .filter(staff | guardians | owner)
            .exclude(pk=self.request.user.pk)
            .exclude(is_archived=True)
            .exclude(role=CustomUser.Role.STUDENT)
            .distinct()
        )


class ConversationDetailView(CommunicationMixin, APIView):
    """Messages d'un fil, et réponse."""

    def get_conversation(self, pk):
        try:
            return Conversation.objects.prefetch_related("participants").get(
                pk=pk, school=self.get_school(), participants=self.request.user,
            )
        except (Conversation.DoesNotExist, ValueError):
            raise serializers.ValidationError(
                {"conversation": "Conversation introuvable."}
            )

    def get(self, request, school_pk, pk):
        conversation = self.get_conversation(pk)
        messages = conversation.messages.select_related("sender").prefetch_related("attachments")

        # Ouvrir le fil vaut lecture : on ne marque que ce qui vient d'autrui.
        unread = messages.exclude(sender=request.user).exclude(read_by=request.user)
        for message in unread:
            message.read_by.add(request.user)

        return Response({
            "id": conversation.id,
            "subject": conversation.subject,
            "participants": [
                {
                    "id": person.id,
                    "name": person.get_full_name(),
                    "role": person.get_role_display(),
                }
                for person in conversation.participants.all()
            ],
            "messages": MessageSerializer(
                messages, many=True, context={"request": request},
            ).data,
        })

    @transaction.atomic
    def post(self, request, school_pk, pk):
        conversation = self.get_conversation(pk)
        body = str(request.data.get("body", "")).strip()
        files = request.FILES.getlist("attachments")
        # Une photo ou un vocal se suffit à lui-même ; c'est le message sans
        # texte *ni* fichier qui n'a pas de sens.
        if not body and not files:
            raise serializers.ValidationError({"body": "Le message ne peut pas être vide."})

        message = Message.objects.create(
            conversation=conversation, sender=request.user, body=body,
        )
        saved = save_attachments(message, files)
        duration = attachment_duration(request)
        if duration and len(saved) == 1:
            saved[0].duration_seconds = duration
            saved[0].save(update_fields=["duration_seconds"])

        message.read_by.add(request.user)
        conversation.last_message_at = message.sent_at
        conversation.save(update_fields=["last_message_at"])
        return Response(
            MessageSerializer(message, context={"request": request}).data,
            status=201,
        )


class MessageRecipientsView(ConversationListView):
    """Personnes à qui l'utilisateur peut écrire dans cette école."""

    def get(self, request, school_pk):
        people = self.reachable_users().order_by("last_name", "first_name")
        return Response({
            "recipients": [
                {
                    "id": person.id,
                    "name": person.get_full_name(),
                    "role": person.get_role_display(),
                    "role_value": person.role,
                }
                for person in people[:500]
            ],
        })

    def post(self, request, school_pk):
        # Cette vue n'hérite de `ConversationListView` que pour `reachable_users` ;
        # créer un fil passe par la liste des conversations.
        raise serializers.ValidationError(
            {"detail": "Utilisez la liste des conversations pour ouvrir un fil."}
        )
